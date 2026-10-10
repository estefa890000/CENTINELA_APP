import streamlit as st
import streamlit.components.v1 as components
import numpy as np
import pandas as pd
import joblib
import base64
import time
import os
import re
import random
import io
import json
import html
import sqlite3
import hashlib
import hmac
from contextlib import closing
from datetime import datetime, date, timedelta
import plotly.graph_objects as go

try:
    import phishing as ph
except ImportError:
    ph = None

try:
    import correos as cx
except ImportError:
    cx = None

try:
    import conector as imap_conector
except ImportError:
    imap_conector = None

try:
    import asistente
except ImportError:
    asistente = None

# =====================================================================
# CONFIGURACION  (cambia la clave del administrador por una tuya)
# =====================================================================
ADMIN_CORREO = "admin@centinela.com"
ADMIN_CLAVE = "Centinela2026!"
DB_PATH = os.environ.get("CENTINELA_DB", "centinela_usuarios.db")

NOMBRES_AMENAZA = {
    "Normal": "Tráfico normal",
    "DoS": "Denegación de servicio (DoS)",
    "Probe": "Escaneo de red (Probe)",
    "R2L": "Acceso no autorizado (R2L)",
    "U2R": "Escalada de privilegios (U2R)",
}
COLOR_CAT = {
    "Normal": "#6B7B8C", "DoS": "#E8843C", "Probe": "#E0A93B",
    "R2L": "#E23D3D", "U2R": "#8E1B1B",
}
SEVERIDAD = {"Normal": "BAJO", "DoS": "ALTO", "Probe": "MEDIO", "R2L": "CRÍTICO", "U2R": "CRÍTICO"}
SEV_ESTILO = {
    "CRÍTICO": ("#C62828", "#FDE8E8"), "ALTO": ("#C2410C", "#FFEBDD"),
    "MEDIO": ("#946200", "#FFF4D6"), "BAJO": ("#2E7D32", "#E3F4E5"),
}
ACTIVO_SIMULADO = {
    "DoS": "Servidor Web ERP", "Probe": "Gateway Principal",
    "R2L": "Servidor Auth LDAP", "U2R": "Estación Contabilidad",
}
CATEGORIAS = ["DoS", "Probe", "R2L", "U2R"]
ESTADOS = ["Nuevo", "En cuarentena", "En revisión", "Resuelto", "Escalado", "Falso positivo", "Liberado"]
ESTADO_ESTILO = {
    "Nuevo": ("#1D4ED8", "#E6EEFF"), "En revisión": ("#946200", "#FFF4D6"),
    "Resuelto": ("#2E7D32", "#E3F4E5"), "Escalado": ("#C2410C", "#FFEBDD"),
    "Falso positivo": ("#475569", "#E8EDF3"), "En cuarentena": ("#6D28D9", "#EFE9FF"), "Liberado": ("#475569", "#E8EDF3"),
}
COLUMNAS_DETALLE = [
    "duration", "protocol_type", "service", "flag", "src_bytes", "dst_bytes", "count",
    "srv_count", "serror_rate", "num_failed_logins", "logged_in", "num_compromised", "root_shell",
]
DESCRIPCION_CAT = {
    "DoS": "El modelo detectó un patrón de tráfico compatible con un intento de saturar un servicio (denegación de servicio).",
    "Probe": "El modelo detectó un patrón compatible con un escaneo de red: alguien está explorando puertos y servicios expuestos.",
    "R2L": "El modelo detectó un patrón compatible con un intento de acceso remoto no autorizado (por ejemplo, adivinar contraseñas).",
    "U2R": "El modelo detectó un patrón compatible con un intento de escalar privilegios dentro de un equipo (posible malware o ransomware).",
}
ACCIONES_CAT = {
    "DoS": [
        ("Revisar el origen del tráfico", "En el firewall o router identifica quién genera el volumen anormal y bloquéalo temporalmente en el borde de la red."),
        ("Limitar las conexiones", "Configura límites de conexiones por segundo en el firewall o en el servidor para absorber el pico."),
        ("Avisar al proveedor de internet", "Si el ataque continúa, pide filtrado del tráfico aguas arriba."),
    ],
    "Probe": [
        ("Identificar quién escanea", "Revisa el registro del firewall para ubicar el origen del escaneo y bloquéalo si no es legítimo."),
        ("Cerrar lo que no se use", "Desactiva puertos y servicios expuestos que la empresa no necesite."),
        ("Vigilar nuevos intentos", "Un escaneo suele ser el paso previo a un ataque: mantén el monitoreo activo las próximas horas."),
    ],
    "R2L": [
        ("Bloquear temporalmente el origen", "Aplica una regla en el firewall o router para frenar las solicitudes desde ese origen."),
        ("Revisar y cambiar contraseñas", "Fuerza el cambio de contraseña en las cuentas que fueron objetivo y cierra sus sesiones activas."),
        ("Verificar la autenticación de dos factores", "Aunque se descifre la contraseña, el segundo factor impide el acceso."),
    ],
    "U2R": [
        ("Aislar el equipo de la red", "Desconéctalo de la red para evitar que el problema se propague a otros equipos."),
        ("Revisar cuentas con privilegios", "Verifica quién tiene permisos de administrador y cierra las sesiones sospechosas."),
        ("Analizar el equipo y restaurar", "Pásale un antivirus o EDR; si hay señales de malware o ransomware, restaura desde una copia de respaldo."),
    ],
}
NOMBRES_AMENAZA["Phishing"] = "Correo de phishing"
COLOR_CAT["Phishing"] = "#7B61FF"
CATEGORIAS_TODAS = ["Phishing", "DoS", "Probe", "R2L", "U2R"]
DESCRIPCION_CAT["Phishing"] = ("Centinela analizó este correo y encontró señales de phishing: un intento de engañar a quien lo recibe "
                               "para que entregue datos o dinero, o para que abra algo peligroso.")
ACCIONES_CAT["Phishing"] = [
    ("No respondas ni abras nada", "No hagas clic en los enlaces ni abras los adjuntos. Si alguien ya lo hizo, avísale al equipo de TI de inmediato."),
    ("Verifica por otro canal", "Confirma con la entidad o la persona que supuestamente lo envió: llama a un número oficial o escribe tú la dirección en el navegador."),
    ("Cambia claves y reporta", "Si se ingresaron datos, cambia las contraseñas y activa la verificación en dos pasos. Reporta el correo como phishing en tu proveedor para bloquear al remitente."),
]
MAX_CORREOS = 10
MAX_BYTES_CORREO = 5_000_000
ARCHIVO_SENSOR = "sensor_eventos.csv"
UMBRAL_ENLACE_MEDIO = 0.30
UMBRAL_ENLACE_ALTO = 0.70
MAX_ENLACES = 20
EJEMPLO_SOSPECHOSO = ("http://bancolombia-seguro-login.xyz/verificar/cuenta?redirect=http://evil.ru/x\n"
                      "http://nequi-actualiza-datos.com/login.php\n192.0.2.10/dian/devolucion.php")
EJEMPLO_NORMAL = "https://www.bancolombia.com/personas\nhttps://www.dian.gov.co/\nhttps://www.amigo.edu.co/"
PROB_RAFAGA = {"Baja": 0.02, "Media": 0.06, "Alta": 0.14}
LONG_RAFAGA = {"DoS": (6, 14), "Probe": (3, 8), "R2L": (1, 3), "U2R": (1, 2)}
VELOCIDAD_SENSOR = {"Lenta": 2, "Media": 4, "Rápida": 8}
COLUMNAS_MODELO = [
    "duration","protocol_type","service","flag","src_bytes","dst_bytes","land",
    "wrong_fragment","urgent","hot","num_failed_logins","logged_in",
    "num_compromised","root_shell","su_attempted","num_root","num_file_creations",
    "num_shells","num_access_files","num_outbound_cmds","is_host_login",
    "is_guest_login","count","srv_count","serror_rate","srv_serror_rate",
    "rerror_rate","srv_rerror_rate","same_srv_rate","diff_srv_rate",
    "srv_diff_host_rate","dst_host_count","dst_host_srv_count",
    "dst_host_same_srv_rate","dst_host_diff_srv_rate","dst_host_same_src_port_rate",
    "dst_host_srv_diff_host_rate","dst_host_serror_rate","dst_host_srv_serror_rate",
    "dst_host_rerror_rate","dst_host_srv_rerror_rate"
]


# =====================================================================
# BASE DE DATOS DE USUARIOS (SQLite, un archivo local)
# =====================================================================
def ahora():
    return datetime.now().strftime("%d/%m/%Y %H:%M")


def hash_clave(clave):
    sal = os.urandom(16)
    h = hashlib.pbkdf2_hmac("sha256", clave.encode(), sal, 100_000)
    return sal.hex() + ":" + h.hex()


def verificar_clave(clave, guardado):
    try:
        sal_hex, h_hex = guardado.split(":")
        h = hashlib.pbkdf2_hmac("sha256", clave.encode(), bytes.fromhex(sal_hex), 100_000)
        return hmac.compare_digest(h.hex(), h_hex)
    except Exception:
        return False


def conectar():
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    return con


def init_db():
    with closing(conectar()) as con:
        con.execute("""CREATE TABLE IF NOT EXISTS usuarios (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            empresa TEXT NOT NULL,
            correo TEXT NOT NULL UNIQUE,
            clave TEXT NOT NULL,
            rol TEXT NOT NULL DEFAULT 'empresa',
            estado TEXT NOT NULL DEFAULT 'pendiente',
            fecha TEXT NOT NULL)""")
        con.execute("""CREATE TABLE IF NOT EXISTS analisis (
            id INTEGER PRIMARY KEY AUTOINCREMENT, usuario_id INTEGER NOT NULL, fecha TEXT NOT NULL,
            total INTEGER NOT NULL, n_normal INTEGER NOT NULL, n_dos INTEGER NOT NULL,
            n_probe INTEGER NOT NULL, n_r2l INTEGER NOT NULL, n_u2r INTEGER NOT NULL)""")
        con.execute("""CREATE TABLE IF NOT EXISTS incidentes (
            id INTEGER PRIMARY KEY AUTOINCREMENT, analisis_id INTEGER NOT NULL, usuario_id INTEGER NOT NULL,
            fecha TEXT NOT NULL, categoria TEXT NOT NULL, severidad TEXT NOT NULL, activo TEXT NOT NULL,
            protocolo TEXT, evento_num INTEGER, certeza REAL, detalle TEXT,
            estado TEXT NOT NULL DEFAULT 'Nuevo')""")
        con.execute("""CREATE TABLE IF NOT EXISTS incidente_historial (
            id INTEGER PRIMARY KEY AUTOINCREMENT, incidente_id INTEGER NOT NULL, fecha TEXT NOT NULL,
            autor TEXT NOT NULL, tipo TEXT NOT NULL, texto TEXT NOT NULL)""")
        for tabla in ("analisis", "incidentes"):
            columnas = [r["name"] for r in con.execute("PRAGMA table_info(" + tabla + ")").fetchall()]
            if "origen" not in columnas:
                con.execute("ALTER TABLE " + tabla + " ADD COLUMN origen TEXT NOT NULL DEFAULT 'csv'")
        if "n_phishing" not in [r["name"] for r in con.execute("PRAGMA table_info(analisis)").fetchall()]:
            con.execute("ALTER TABLE analisis ADD COLUMN n_phishing INTEGER NOT NULL DEFAULT 0")
        con.execute("""CREATE TABLE IF NOT EXISTS correos_procesados (
            id INTEGER PRIMARY KEY AUTOINCREMENT, usuario_id INTEGER NOT NULL, mensaje_id TEXT NOT NULL, fecha TEXT NOT NULL,
            UNIQUE(usuario_id, mensaje_id))""")
        adm = con.execute("SELECT * FROM usuarios WHERE correo=?", (ADMIN_CORREO,)).fetchone()
        if adm is None:
            con.execute(
                "INSERT INTO usuarios (empresa, correo, clave, rol, estado, fecha) VALUES (?,?,?,?,?,?)",
                ("Centinela · Administración", ADMIN_CORREO, hash_clave(ADMIN_CLAVE), "admin", "aprobado", ahora()))
        elif not verificar_clave(ADMIN_CLAVE, adm["clave"]):
            con.execute("UPDATE usuarios SET clave=? WHERE correo=?", (hash_clave(ADMIN_CLAVE), ADMIN_CORREO))
        con.commit()


def crear_solicitud(empresa, correo, clave, clave2):
    empresa, correo = empresa.strip(), correo.strip().lower()
    if not empresa or not correo or not clave:
        return False, "Completa todos los campos."
    if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", correo):
        return False, "Escribe un correo válido."
    if len(clave) < 8:
        return False, "La contraseña debe tener mínimo 8 caracteres."
    if clave != clave2:
        return False, "Las contraseñas no coinciden."
    try:
        with closing(conectar()) as con:
            con.execute(
                "INSERT INTO usuarios (empresa, correo, clave, rol, estado, fecha) VALUES (?,?,?,?,?,?)",
                (empresa, correo, hash_clave(clave), "empresa", "pendiente", ahora()))
            con.commit()
    except sqlite3.IntegrityError:
        return False, "Ese correo ya está registrado o ya tiene una solicitud en curso."
    return True, "ok"


def autenticar(correo, clave):
    correo = correo.strip().lower()
    with closing(conectar()) as con:
        fila = con.execute("SELECT * FROM usuarios WHERE correo=?", (correo,)).fetchone()
    if fila is None or not verificar_clave(clave, fila["clave"]):
        return "invalido", None
    return fila["estado"], dict(fila)


def cambiar_clave(user_id, actual, nueva, nueva2):
    with closing(conectar()) as con:
        fila = con.execute("SELECT * FROM usuarios WHERE id=?", (user_id,)).fetchone()
        if fila is None or not verificar_clave(actual, fila["clave"]):
            return False, "La contraseña actual no es correcta."
        if len(nueva) < 8:
            return False, "La nueva contraseña debe tener mínimo 8 caracteres."
        if nueva != nueva2:
            return False, "Las contraseñas nuevas no coinciden."
        if nueva == actual:
            return False, "La nueva contraseña debe ser diferente a la actual."
        con.execute("UPDATE usuarios SET clave=? WHERE id=?", (hash_clave(nueva), user_id))
        con.commit()
    return True, "Contraseña actualizada. Úsala la próxima vez que inicies sesión."


def cambiar_empresa(user_id, nombre):
    nombre = (nombre or "").strip()
    if len(nombre) < 2:
        return False, "Escribe el nombre de la empresa."
    with closing(conectar()) as con:
        con.execute("UPDATE usuarios SET empresa=? WHERE id=?", (nombre[:120], user_id))
        con.commit()
    return True, "Nombre de la empresa actualizado."


def datos_cuenta(user_id):
    with closing(conectar()) as con:
        fila = con.execute("SELECT id, empresa, correo, rol, estado, fecha FROM usuarios WHERE id=?", (user_id,)).fetchone()
        n_inc = con.execute("SELECT COUNT(*) FROM incidentes WHERE usuario_id=?", (user_id,)).fetchone()[0]
        n_corr = con.execute("SELECT COUNT(*) FROM correos_procesados WHERE usuario_id=?", (user_id,)).fetchone()[0]
    return (dict(fila) if fila else {}), n_inc, n_corr


def contar_pendientes():
    with closing(conectar()) as con:
        return con.execute("SELECT COUNT(*) FROM usuarios WHERE estado='pendiente'").fetchone()[0]


def listar_usuarios(estado):
    with closing(conectar()) as con:
        return [dict(f) for f in con.execute(
            "SELECT id, empresa, correo, fecha, estado FROM usuarios WHERE estado=? AND rol='empresa' ORDER BY id DESC",
            (estado,)).fetchall()]


def cambiar_estado(user_id, estado):
    with closing(conectar()) as con:
        con.execute("UPDATE usuarios SET estado=? WHERE id=? AND rol='empresa'", (estado, user_id))
        con.commit()


# ---------------------------------------------------------------------
# ANALISIS E INCIDENTES (historial persistente)
# ---------------------------------------------------------------------
def ahora_iso():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def fmt_fecha(iso, corto=False):
    try:
        d = datetime.strptime(iso, "%Y-%m-%d %H:%M:%S")
        return d.strftime("%H:%M:%S") if corto else d.strftime("%d/%m/%Y %H:%M:%S")
    except Exception:
        return str(iso)


def _alcance(usuario):
    if usuario["rol"] == "admin":
        return "", []
    return " AND usuario_id=?", [usuario["id"]]


def guardar_analisis(usuario_id, conteo, total, incidentes):
    ts = ahora_iso()
    with closing(conectar()) as con:
        cur = con.execute(
            "INSERT INTO analisis (usuario_id, fecha, total, n_normal, n_dos, n_probe, n_r2l, n_u2r) VALUES (?,?,?,?,?,?,?,?)",
            (usuario_id, ts, total, conteo["Normal"], conteo["DoS"], conteo["Probe"], conteo["R2L"], conteo["U2R"]))
        aid = cur.lastrowid
        for inc in incidentes:
            c2 = con.execute(
                "INSERT INTO incidentes (analisis_id, usuario_id, fecha, categoria, severidad, activo, protocolo, evento_num, certeza, detalle, estado) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                (aid, usuario_id, inc["fecha"], inc["categoria"], inc["severidad"], inc["activo"],
                 inc["protocolo"], inc["evento_num"], inc["certeza"], inc["detalle"], "Nuevo"))
            con.execute(
                "INSERT INTO incidente_historial (incidente_id, fecha, autor, tipo, texto) VALUES (?,?,?,?,?)",
                (c2.lastrowid, inc["fecha"], "Centinela", "sistema",
                 "Alerta generada por el modelo (certeza {:.0f}%).".format(inc["certeza"] * 100)))
        con.commit()
    return aid, ts


def iniciar_sesion_sensor(usuario_id):
    with closing(conectar()) as con:
        cur = con.execute(
            "INSERT INTO analisis (usuario_id, fecha, total, n_normal, n_dos, n_probe, n_r2l, n_u2r, origen) VALUES (?,?,0,0,0,0,0,0,'sensor')",
            (usuario_id, ahora_iso()))
        con.commit()
        return cur.lastrowid


def registrar_lote_sensor(sesion_id, usuario_id, conteo_lote, incidentes):
    with closing(conectar()) as con:
        con.execute(
            "UPDATE analisis SET total=total+?, n_normal=n_normal+?, n_dos=n_dos+?, n_probe=n_probe+?, n_r2l=n_r2l+?, n_u2r=n_u2r+? "
            "WHERE id=? AND usuario_id=?",
            (sum(conteo_lote.values()), conteo_lote["Normal"], conteo_lote["DoS"], conteo_lote["Probe"],
             conteo_lote["R2L"], conteo_lote["U2R"], sesion_id, usuario_id))
        for inc in incidentes:
            c2 = con.execute(
                "INSERT INTO incidentes (analisis_id, usuario_id, fecha, categoria, severidad, activo, protocolo, evento_num, certeza, detalle, estado, origen) "
                "VALUES (?,?,?,?,?,?,?,?,?,?,?,'sensor')",
                (sesion_id, usuario_id, inc["fecha"], inc["categoria"], inc["severidad"], inc["activo"],
                 inc["protocolo"], inc["evento_num"], inc["certeza"], inc["detalle"], "Nuevo"))
            con.execute(
                "INSERT INTO incidente_historial (incidente_id, fecha, autor, tipo, texto) VALUES (?,?,?,?,?)",
                (c2.lastrowid, inc["fecha"], "Centinela", "sistema",
                 "Alerta generada automáticamente por el sensor virtual (certeza {:.0f}%).".format(inc["certeza"] * 100)))
        con.commit()


# ---------------------------------------------------------------------
# SENSOR VIRTUAL (simula trafico de red que Centinela analiza solo)
# ---------------------------------------------------------------------
CLASES = ["Normal", "DoS", "Probe", "R2L", "U2R"]


def nuevo_estado_sensor():
    return {
        "activo": False, "sesion_id": None, "tick": 0, "ultimo_tick": 0.0, "ultimo_rerun": 0.0, "ultimo_toast": 0.0,
        "total": 0, "conteo": {c: 0 for c in CLASES},
        "verdad": {c: {"enviados": 0, "alertados": 0, "exactos": 0} for c in CLASES},
        "feed": [], "serie": [], "rafaga": None, "alertas_sesion": 0, "criticas_sesion": 0,
    }


def generar_lote(estado, grupos, intensidad, max_eventos, tipos, rng=random):
    filas = []
    for _ in range(rng.randint(1, max_eventos)):
        if estado["rafaga"] is None and tipos and rng.random() < PROB_RAFAGA[intensidad]:
            tipo = rng.choice(tipos)
            lo, hi = LONG_RAFAGA[tipo]
            estado["rafaga"] = {"tipo": tipo, "restantes": rng.randint(lo, hi)}
        if estado["rafaga"] is not None:
            tipo = estado["rafaga"]["tipo"]
            estado["rafaga"]["restantes"] -= 1
            if estado["rafaga"]["restantes"] <= 0:
                estado["rafaga"] = None
        else:
            tipo = "Normal"
        grupo = grupos.get(tipo)
        if grupo is None or len(grupo) == 0:
            grupo = grupos["Normal"]
        filas.append(grupo.iloc[rng.randrange(len(grupo))])
    return pd.DataFrame(filas).reset_index(drop=True)


def aplicar_lote(estado, df_lote, cats, certezas):
    incidentes = []
    conteo_lote = {c: 0 for c in CLASES}
    for i, cat in enumerate(cats):
        fila = df_lote.iloc[i]
        real = str(fila["etiqueta"])
        estado["total"] += 1
        estado["conteo"][cat] += 1
        conteo_lote[cat] += 1
        v = estado["verdad"][real]
        v["enviados"] += 1
        if cat != "Normal":
            v["alertados"] += 1
        if cat == real:
            v["exactos"] += 1
        estado["feed"].append({"n": estado["total"], "hora": datetime.now().strftime("%H:%M:%S"), "cat": cat,
                               "certeza": float(certezas[i]), "proto": str(fila["protocol_type"])})
        if cat != "Normal":
            det = {c: (fila[c].item() if hasattr(fila[c], "item") else fila[c]) for c in COLUMNAS_DETALLE}
            incidentes.append({
                "fecha": ahora_iso(), "categoria": cat, "severidad": SEVERIDAD[cat], "activo": ACTIVO_SIMULADO[cat],
                "protocolo": str(fila["protocol_type"]), "evento_num": estado["total"], "certeza": float(certezas[i]),
                "detalle": json.dumps(det, default=str),
            })
    estado["feed"] = estado["feed"][-14:]
    estado["serie"] = (estado["serie"] + [len(incidentes)])[-60:]
    estado["alertas_sesion"] += len(incidentes)
    estado["criticas_sesion"] += sum(1 for i in incidentes if i["severidad"] == "CRÍTICO")
    estado["tick"] += 1
    return incidentes, conteo_lote


@st.cache_data
def cargar_grupos_sensor():
    df = pd.read_csv(ARCHIVO_SENSOR)
    return {c: g.reset_index(drop=True) for c, g in df.groupby("etiqueta")}


def guardar_analisis_correos(usuario_id, resultados, origen="correo"):
    ts = ahora_iso()
    graves = [r for r in resultados if r["nivel"] in ("ALTO", "MEDIO")]
    ids = []
    with closing(conectar()) as con:
        cur = con.execute(
            "INSERT INTO analisis (usuario_id, fecha, total, n_normal, n_dos, n_probe, n_r2l, n_u2r, origen, n_phishing) VALUES (?,?,?,?,0,0,0,0,?,?)",
            (usuario_id, ts, len(resultados), len(resultados) - len(graves), origen, len(graves)))
        aid = cur.lastrowid
        for k, r in enumerate(graves, start=1):
            def pct(x):
                return "{:.0f}%".format(x * 100) if x is not None else "no evaluado"
            det = {"remitente": (r["de_correo"] or r["de_nombre"] or "desconocido")[:120], "asunto": r["asunto"][:140], "idioma": r["idioma"],
                   "SPF": r["auth"]["spf"] or "sin dato", "DKIM": r["auth"]["dkim"] or "sin dato", "DMARC": r["auth"]["dmarc"] or "sin dato",
                   "riesgo_del_texto": pct(r["p_texto"]), "riesgo_de_enlaces": pct(r["p_enlaces"]), "riesgo_por_senales_tecnicas": pct(r["p_tecnico"]),
                   "enlaces_encontrados": len(r["enlaces"]), "adjuntos": len(r["adjuntos"]),
                   "razones": "; ".join(x["texto"] for x in r["razones"][:6]),
                   "mensaje_id": (r.get("mensaje_id") or "")[:300], "ubicacion": "cuarentena" if r.get("en_cuarentena") else "bandeja"}
            c2 = con.execute(
                "INSERT INTO incidentes (analisis_id, usuario_id, fecha, categoria, severidad, activo, protocolo, evento_num, certeza, detalle, estado, origen) "
                "VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                (aid, usuario_id, ts, "Phishing", r["nivel"], "Buzón de correo", "correo", k, float(r["p_final"]), json.dumps(det, ensure_ascii=False),
                 "En cuarentena" if r.get("en_cuarentena") else "Nuevo", origen))
            con.execute("INSERT INTO incidente_historial (incidente_id, fecha, autor, tipo, texto) VALUES (?,?,?,?,?)",
                        (c2.lastrowid, ts, "Centinela", "sistema", "Correo marcado por Centinela con riesgo " + r["nivel"].lower() + " (" + pct(r["p_final"]) + ")."
                         + (" Se movió a la carpeta de cuarentena del buzón." if r.get("en_cuarentena") else "")))
            ids.append(c2.lastrowid)
        con.commit()
    return aid, ids


def marcar_procesado(usuario_id, mensaje_id):
    """True si el correo era nuevo (y queda marcado); False si ya se había procesado."""
    with closing(conectar()) as con:
        cur = con.execute("INSERT OR IGNORE INTO correos_procesados (usuario_id, mensaje_id, fecha) VALUES (?,?,?)", (usuario_id, mensaje_id, ahora_iso()))
        con.commit()
        return cur.rowcount > 0


def listar_incidentes(usuario, analisis_id=None, desde=None, hasta=None, cats=None, estados=None):
    if cats is not None and len(cats) == 0:
        return []
    sql, params = "SELECT * FROM incidentes WHERE 1=1", []
    extra, p = _alcance(usuario)
    sql += extra
    params += p
    if analisis_id is not None:
        sql += " AND analisis_id=?"
        params.append(analisis_id)
    if desde:
        sql += " AND fecha>=?"
        params.append(str(desde) + " 00:00:00")
    if hasta:
        sql += " AND fecha<=?"
        params.append(str(hasta) + " 23:59:59")
    if cats:
        sql += " AND categoria IN (" + ",".join("?" * len(cats)) + ")"
        params += list(cats)
    if estados:
        sql += " AND estado IN (" + ",".join("?" * len(estados)) + ")"
        params += list(estados)
    sql += " ORDER BY id DESC"
    with closing(conectar()) as con:
        return [dict(r) for r in con.execute(sql, params).fetchall()]


def obtener_incidente(usuario, inc_id):
    extra, p = _alcance(usuario)
    with closing(conectar()) as con:
        fila = con.execute("SELECT * FROM incidentes WHERE id=?" + extra, [inc_id] + p).fetchone()
    return dict(fila) if fila else None


def historial_incidente(inc_id):
    with closing(conectar()) as con:
        return [dict(r) for r in con.execute(
            "SELECT * FROM incidente_historial WHERE incidente_id=? ORDER BY id DESC", (inc_id,)).fetchall()]


def cambiar_estado_incidente(usuario, inc_id, estado, autor):
    if estado not in ESTADOS or obtener_incidente(usuario, inc_id) is None:
        return
    with closing(conectar()) as con:
        con.execute("UPDATE incidentes SET estado=? WHERE id=?", (estado, inc_id))
        con.execute(
            "INSERT INTO incidente_historial (incidente_id, fecha, autor, tipo, texto) VALUES (?,?,?,?,?)",
            (inc_id, ahora_iso(), autor, "estado", "Estado cambiado a \u00ab" + estado + "\u00bb por " + autor + "."))
        con.commit()


def agregar_nota(usuario, inc_id, autor, texto):
    if obtener_incidente(usuario, inc_id) is None or not texto.strip():
        return
    with closing(conectar()) as con:
        con.execute(
            "INSERT INTO incidente_historial (incidente_id, fecha, autor, tipo, texto) VALUES (?,?,?,?,?)",
            (inc_id, ahora_iso(), autor, "nota", texto.strip()))
        con.commit()


def contar_nuevos(usuario):
    extra, p = _alcance(usuario)
    with closing(conectar()) as con:
        return con.execute("SELECT COUNT(*) FROM incidentes WHERE estado IN ('Nuevo','En cuarentena')" + extra, p).fetchone()[0]


def ultimo_analisis(usuario):
    extra, p = _alcance(usuario)
    with closing(conectar()) as con:
        fila = con.execute("SELECT fecha FROM analisis WHERE 1=1" + extra + " ORDER BY id DESC LIMIT 1", p).fetchone()
    return fila[0] if fila else None


def conteo_eventos(usuario, desde, hasta):
    extra, p = _alcance(usuario)
    sql = ("SELECT COALESCE(SUM(n_normal),0), COALESCE(SUM(n_dos),0), COALESCE(SUM(n_probe),0), "
           "COALESCE(SUM(n_r2l),0), COALESCE(SUM(n_u2r),0), COALESCE(SUM(n_phishing),0) FROM analisis WHERE fecha>=? AND fecha<=?" + extra)
    with closing(conectar()) as con:
        f = con.execute(sql, [str(desde) + " 00:00:00", str(hasta) + " 23:59:59"] + p).fetchone()
    return {"Normal": f[0], "DoS": f[1], "Probe": f[2], "R2L": f[3], "U2R": f[4], "Phishing": f[5]}


def calcular_reporte(usuario, desde, hasta, cats):
    incs = listar_incidentes(usuario, desde=desde, hasta=hasta, cats=cats)
    ev = conteo_eventos(usuario, desde, hasta)
    dist = {"Normal": ev["Normal"]}
    for c in CATEGORIAS_TODAS:
        if c in cats:
            dist[c] = ev[c]
    tiles = {
        "criticos": sum(1 for i in incs if i["severidad"] == "CRÍTICO"),
        "altos": sum(1 for i in incs if i["severidad"] == "ALTO"),
        "medios": sum(1 for i in incs if i["severidad"] == "MEDIO"),
        "fp": sum(1 for i in incs if i["estado"] == "Falso positivo"),
    }
    return {"incs": incs, "tiles": tiles, "dist": dist, "total": sum(dist.values())}


def _nombre_dist(c):
    return "Sin amenaza detectada" if c == "Normal" else NOMBRES_AMENAZA[c]


def generar_pdf(empresa, autor, desde, hasta, rep):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=1.8 * cm, bottomMargin=1.8 * cm, title="Reporte Centinela")
    est = getSampleStyleSheet()
    oscuro, gris = colors.HexColor("#16202B"), colors.HexColor("#6B7B8C")
    h1 = ParagraphStyle("h1", parent=est["Title"], fontName="Helvetica-Bold", fontSize=20, textColor=oscuro, alignment=0, spaceAfter=2)
    sub = ParagraphStyle("sub", parent=est["Normal"], fontSize=10, textColor=gris, leading=14)
    h2 = ParagraphStyle("h2", parent=est["Heading2"], fontName="Helvetica-Bold", fontSize=12, textColor=oscuro, spaceBefore=14, spaceAfter=6)
    cel = ParagraphStyle("cel", parent=est["Normal"], fontSize=8.5, leading=11)
    pie = ParagraphStyle("pie", parent=est["Normal"], fontSize=8, textColor=gris, leading=11)

    t = rep["tiles"]
    story = [
        Paragraph("CENTINELA - Reporte de seguridad", h1),
        Paragraph(escape(empresa) + "<br/>Periodo: " + desde.strftime("%d/%m/%Y") + " al " + hasta.strftime("%d/%m/%Y"), sub),
        Spacer(1, 12),
    ]
    kpis = Table([
        ["CRITICOS", "ALTOS", "MEDIOS", "FALSOS POSITIVOS"],
        [str(t["criticos"]), str(t["altos"]), str(t["medios"]), str(t["fp"])],
    ], colWidths=[4.25 * cm] * 4)
    kpis.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#FDECEC")), ("BACKGROUND", (1, 0), (1, -1), colors.HexColor("#FFF1E0")),
        ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#FFF8E1")), ("BACKGROUND", (3, 0), (3, -1), colors.HexColor("#E8EDF3")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"), ("FONTSIZE", (0, 1), (-1, 1), 22),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("LINEAFTER", (0, 0), (-2, -1), 4, colors.white),
    ]))
    story += [kpis, Paragraph("Distribucion por categoria", h2)]

    filas = [["Categoria", "Eventos", "%"]]
    for c, n in rep["dist"].items():
        pct = round(n / rep["total"] * 100, 1) if rep["total"] else 0
        filas.append([_nombre_dist(c), str(n), str(pct) + "%"])
    dist = Table(filas, colWidths=[9 * cm, 3.5 * cm, 3.5 * cm])
    dist.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), oscuro), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F7FA")]),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#DDE3EA")),
    ]))
    story += [dist, Paragraph("Incidentes del periodo", h2)]

    if rep["incs"]:
        filas = [["ID", "Fecha", "Tipo de amenaza", "Severidad", "Estado", "Certeza"]]
        for i in rep["incs"][:200]:
            filas.append([str(i["id"]), fmt_fecha(i["fecha"]), Paragraph(escape(NOMBRES_AMENAZA[i["categoria"]]), cel),
                          i["severidad"], i["estado"], "{:.0f}%".format(i["certeza"] * 100)])
        tabla = Table(filas, colWidths=[1.2 * cm, 3.6 * cm, 5.6 * cm, 2.2 * cm, 2.4 * cm, 1.8 * cm], repeatRows=1)
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), oscuro), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F7FA")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#DDE3EA")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        story.append(tabla)
        if len(rep["incs"]) > 200:
            story.append(Paragraph("Se muestran los 200 incidentes mas recientes de " + str(len(rep["incs"])) + ".", pie))
    else:
        story.append(Paragraph("No hay incidentes en el periodo seleccionado.", sub))

    story += [
        Spacer(1, 16),
        Paragraph("Elaborado por " + escape(autor) + " el " + datetime.now().strftime("%d/%m/%Y %H:%M") + ". "
                  "Prototipo academico: el activo afectado es simulado y los eventos provienen de los archivos "
                  "cargados por la empresa. Las cifras de desempeno del modelo (precision, recall, F1) se "
                  "documentan en el informe de validacion del proyecto.", pie),
    ]
    doc.build(story)
    return buf.getvalue()


def generar_excel(empresa, autor, desde, hasta, rep):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = "CENTINELA - Reporte de seguridad"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Empresa"
    ws["B2"] = empresa
    ws["A3"] = "Periodo"
    ws["B3"] = desde.strftime("%d/%m/%Y") + " al " + hasta.strftime("%d/%m/%Y")
    ws["A4"] = "Elaborado por"
    ws["B4"] = autor
    ws["A6"] = "Severidad"
    ws["B6"] = "Incidentes"
    t = rep["tiles"]
    for r, (k, v) in enumerate([("Críticos", t["criticos"]), ("Altos", t["altos"]), ("Medios", t["medios"]),
                                ("Falsos positivos", t["fp"])], start=7):
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=v)
    ws["A12"] = "Categoría"
    ws["B12"] = "Eventos"
    ws["C12"] = "%"
    for r, (c, n) in enumerate(rep["dist"].items(), start=13):
        ws.cell(row=r, column=1, value=_nombre_dist(c))
        ws.cell(row=r, column=2, value=n)
        ws.cell(row=r, column=3, value=round(n / rep["total"] * 100, 1) if rep["total"] else 0)
    cab = PatternFill("solid", fgColor="16202B")
    for celda in ("A6", "B6", "A12", "B12", "C12"):
        ws[celda].font = Font(bold=True, color="FFFFFF")
        ws[celda].fill = cab
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 10

    w2 = wb.create_sheet("Incidentes")
    w2.append(["ID", "Fecha", "Tipo de amenaza", "Activo (simulado)", "Severidad", "Estado", "Certeza"])
    for c in w2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = cab
        c.alignment = Alignment(horizontal="center")
    for i in rep["incs"]:
        w2.append([i["id"], fmt_fecha(i["fecha"]), NOMBRES_AMENAZA[i["categoria"]], i["activo"],
                   i["severidad"], i["estado"], round(i["certeza"] * 100)])
    for col, ancho in zip("ABCDEFG", (8, 22, 32, 26, 12, 16, 10)):
        w2.column_dimensions[col].width = ancho
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# =====================================================================
# MODELO
# =====================================================================
@st.cache_resource
def cargar_modelo():
    return (
        joblib.load("modelo_centinela_multiclase.pkl"),
        joblib.load("target_encoder_centinela.pkl"),
        joblib.load("encoders_centinela.pkl"),
    )


# =====================================================================
# ARRANQUE
# =====================================================================
st.set_page_config(page_title="Centinela", page_icon="isotipo_centinela.png", layout="wide")
init_db()
modelo, target_encoder, encoders = cargar_modelo()


@st.cache_resource
def cargar_phishing():
    import xgboost as xgb
    modelo = xgb.XGBClassifier()
    modelo.load_model("modelo_phishing.json")
    with open("phishing_meta.json", encoding="utf-8") as f:
        meta = json.load(f)
    return modelo, meta


@st.cache_resource
def cargar_correos():
    import xgboost as xgb
    b = xgb.XGBClassifier()
    b.load_model("modelo_enlaces_correo.json")
    return {"es": cx.ModeloTexto("modelo_texto_es.json"), "en": cx.ModeloTexto("modelo_texto_en.json"), "urls": cx.fabrica_puntuar_urls(b)}


def nivel_enlace(p):
    if p >= UMBRAL_ENLACE_ALTO:
        return "ALTO"
    return "MEDIO" if p >= UMBRAL_ENLACE_MEDIO else "BAJO"


def validar_enlace(texto):
    t = texto.strip()
    if re.fullmatch(r"[^@\s/]+@[^@\s/]+\.[^@\s/]+", t):
        return None, "Esto parece una dirección de correo, no un enlace. Para revisar un correo usa «Analizar correos»."
    if len(t) > 2000:
        return None, "El enlace es demasiado largo (máximo 2000 caracteres)."
    if re.search(r"\s", t):
        return None, "Un enlace no lleva espacios. Revisa que sea un solo enlace por línea."
    if "." not in t and ":" not in t:
        return None, "No parece un enlace (le falta el punto del dominio, por ejemplo sitio.com)."
    return t, None


def analizar_enlaces(lineas, modelo):
    resultados, validos, posiciones = [], [], []
    for linea in lineas:
        t, error = validar_enlace(linea)
        resultados.append({"url": linea.strip(), "error": error})
        if t:
            validos.append(t)
            posiciones.append(len(resultados) - 1)
    if validos:
        X = np.array([ph.vector(u) for u in validos], dtype=np.float32)
        probas = modelo.predict_proba(X)[:, 1]
        for k, u, pr in zip(posiciones, validos, probas):
            razones = ph.razones(u)
            if ph.extraer(u)["es_acortador"]:
                resultados[k].update({"nivel": "MEDIO", "p": None, "acortado": True, "razones": razones})
            else:
                resultados[k].update({"nivel": nivel_enlace(float(pr)), "p": float(pr), "acortado": False, "razones": razones})
    return resultados


def clasificar_df(df):
    X = df[COLUMNAS_MODELO].copy()
    for col, encoder in encoders.items():
        X[col] = X[col].map(lambda v: v if v in encoder.classes_ else encoder.classes_[0])
        X[col] = encoder.transform(X[col])
    proba = modelo.predict_proba(X)
    idx = proba.argmax(axis=1)
    cats = list(target_encoder.inverse_transform(modelo.classes_[idx]))
    certezas = [float(proba[i, j]) for i, j in enumerate(idx)]
    return cats, certezas

if "pagina" not in st.session_state:
    st.session_state["pagina"] = "login"
if st.session_state["pagina"] == "dashboard" and "usuario" not in st.session_state:
    st.session_state["pagina"] = "login"
pagina = st.session_state["pagina"]

with open("isotipo_centinela.png", "rb") as f:
    logo_b64 = base64.b64encode(f.read()).decode()

if pagina in ("login", "registro"):
    fondo_css = """
    html, body, [data-testid="stAppViewContainer"], .stApp {
        background-color: #0B111A !important;
        background-image:
            repeating-linear-gradient(60deg, rgba(46,124,228,0.14) 0 1px, transparent 1px 52px),
            repeating-linear-gradient(-60deg, rgba(46,124,228,0.14) 0 1px, transparent 1px 52px),
            repeating-linear-gradient(0deg, rgba(46,124,228,0.14) 0 1px, transparent 1px 90px),
            radial-gradient(circle at 15% 15%, rgba(46,124,228,0.22), transparent 45%),
            radial-gradient(circle at 85% 80%, rgba(60,140,64,0.16), transparent 45%),
            linear-gradient(180deg, #0B111A 0%, #16202B 60%, #1B2A3D 100%) !important;
        background-repeat: repeat, repeat, repeat, no-repeat, no-repeat, no-repeat !important;
        background-position: center, center, center, left top, right bottom, center !important;
        background-attachment: fixed !important;
    }
    [data-testid="stSidebar"], [data-testid="stSidebarCollapsedControl"] { display: none !important; }
    .block-container { max-width: 520px; margin: 2.5rem auto 0 auto; background-color: #FFFFFF;
        border-radius: 22px; padding: 2.2rem 2rem 1.8rem 2rem; box-shadow: 0 20px 50px rgba(0,0,0,0.45); }
    """
else:
    fondo_css = """
    html, body, [data-testid="stAppViewContainer"], .stApp { background-color: #F4F7FA !important; }
    .block-container { max-width: 100%; padding: 3.2rem 2.2rem 2rem 2.2rem; }
    [data-testid="stSidebar"] { background-color: #16202B !important; width: 250px !important; min-width: 250px !important; }
    [data-testid="stSidebar"] * { color: #FFFFFF; }
    """

CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');
html, body, .stApp, .stMarkdown, .stButton button, input, label, p { font-family: 'Inter', sans-serif; }
__FONDO__
[data-testid="stAppDeployButton"], [data-testid="stMainMenu"] { display: none !important; }
[data-testid="stHeader"] { background: transparent !important; }
.block-container, .block-container p, .block-container span, .block-container label { color: #16202B; }

/* ---------- login / registro ---------- */
.logo-box { width: 72px; height: 72px; border-radius: 18px; background: #EAF2FE;
    display: flex; align-items: center; justify-content: center; margin: 0 auto 14px auto; }
.logo-box img { width: 40px; }
.card-title { text-align: center; font-size: 28px; font-weight: 800; letter-spacing: .5px; margin: 0; }
.card-subtitle { text-align: center; font-size: 15px; font-weight: 700; color: #2E7CE4; margin: 2px 0 6px 0; }
.card-desc { text-align: center; font-size: 13px; color: #6B7B8C; margin-bottom: 18px; line-height: 1.45; }
.field-label { font-size: 13px; font-weight: 700; margin: 12px 0 4px 0; }
.divider-text { text-align: center; font-size: 11px; color: #8FA0B3; letter-spacing: 2px; margin: 16px 0 12px 0; }
.footer-sub { text-align: center; font-size: 12px; color: #6B7B8C; margin-top: 14px; }
.footer-mono { text-align: center; font-family: Consolas, monospace; font-size: 11px; color: #6B7B8C;
    background: #F1F5F9; border-radius: 8px; padding: 6px 10px; margin: 8px auto 0 auto; width: fit-content; }

div[data-testid="stTextInput"] label { display: none !important; }
div[data-testid="stTextInput"] input {
    background-color: #F8FAFC !important; color: #16202B !important;
    border: 1px solid #E2E8F0 !important; border-radius: 12px !important; height: 46px;
    padding-left: 40px !important; background-repeat: no-repeat; background-position: 13px center; background-size: 17px;
}
div[data-testid="stTextInput"] input:focus { border-color: #2E7CE4 !important; box-shadow: 0 0 0 3px rgba(46,124,228,.15) !important; }
div[data-testid="stTextInput"]:has(input[aria-label="correo_input"]) input,
div[data-testid="stTextInput"]:has(input[aria-label="reg_correo"]) input {
    background-image: url("data:image/svg+xml;utf8,<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 24 24' fill='none' stroke='%236B7B8C' stroke-width='2'><rect x='3' y='5' width='18' height='14' rx='2'/><path d='M3 7l9 6 9-6'/></svg>");
}
div[data-testid="stTextInput"]:has(input[aria-label="clave_input"]) input,
div[data-testid="stTextInput"]:has(input[aria-label="reg_clave"]) input,
div[data-testid="stTextInput"]:has(input[aria-label="reg_clave2"]) input {
    background-image: url("data:image/svg+xml;utf8,<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 24 24' fill='none' stroke='%236B7B8C' stroke-width='2'><rect x='4' y='10' width='16' height='10' rx='2'/><path d='M8 10V7a4 4 0 0 1 8 0v3'/></svg>");
}
div[data-testid="stTextInput"]:has(input[aria-label="reg_empresa"]) input {
    background-image: url("data:image/svg+xml;utf8,<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 24 24' fill='none' stroke='%236B7B8C' stroke-width='2'><rect x='4' y='3' width='16' height='18' rx='2'/><path d='M9 8h2M13 8h2M9 12h2M13 12h2M10 21v-4h4v4'/></svg>");
}
.stButton button[kind="primary"] { background: #2E7CE4; border: none; border-radius: 12px; height: 46px; font-weight: 700; color: #fff; }
.stButton button[kind="primary"]:hover { background: #2569c4; }
.stButton button[kind="secondary"] { border-radius: 12px; font-weight: 600; }
.st-key-forgot button, .st-key-ir_registro button, .st-key-ir_login button {
    color: #2E7CE4 !important; font-weight: 700 !important; padding: 0 !important; min-height: 0 !important; }
.st-key-footer_reg { background: #EEF4FF; border-radius: 12px; padding: 12px 16px; }
.st-key-footer_reg p { margin: 0; font-weight: 600; font-size: 14px; }

/* ---------- dashboard ---------- */
[class*="st-key-card_"] { background: #FFFFFF; border-radius: 16px; padding: 18px 20px;
    box-shadow: 0 2px 12px rgba(22,32,43,0.06); margin-bottom: 14px; }
.topbar-card { background:#FFFFFF; border-radius:16px; padding:14px 22px; display:flex; align-items:center;
    justify-content:space-between; box-shadow:0 2px 12px rgba(22,32,43,0.06); margin-bottom:18px; }
.company-name { font-size:16px; font-weight:800; }
.company-sub { font-size:12px; color:#6B7B8C; }
.topbar-right { display:flex; align-items:center; gap:18px; }
.badge-live { background:#E3F4E5; color:#2E7D32; font-size:12px; font-weight:700; padding:6px 12px; border-radius:20px; display:flex; align-items:center; gap:6px; }
.badge-live .dot { width:8px; height:8px; border-radius:50%; background:#3C8C40; display:inline-block; }
.bell { position:relative; width:22px; height:22px; }
.bell-badge { position:absolute; top:-8px; right:-9px; background:#E23D3D; color:#fff !important; font-size:10px; font-weight:800;
    min-width:17px; height:17px; border-radius:9px; display:flex; align-items:center; justify-content:center; padding:0 4px; }
.user-chip { display:flex; align-items:center; gap:8px; font-size:13px; font-weight:600; }
.avatar { width:34px; height:34px; border-radius:50%; background:linear-gradient(135deg,#2E7CE4,#3C8C40); color:#fff !important;
    display:flex; align-items:center; justify-content:center; font-size:12px; font-weight:800; }

.dash-head { display:flex; align-items:center; justify-content:space-between; margin-bottom:16px; }
.dash-title { font-size:24px; font-weight:800; display:flex; align-items:center; gap:10px; }
.tag-pyme { background:#E6EEFF; color:#1D4ED8 !important; font-size:11px; font-weight:800; padding:3px 9px; border-radius:6px; }
.dash-sub { font-size:13px; color:#6B7B8C; margin-top:2px; }
.last-box { background:#fff; border-radius:12px; padding:8px 16px; box-shadow:0 2px 12px rgba(22,32,43,0.06); text-align:right; }
.last-label { font-size:11px; color:#6B7B8C; font-weight:600; }
.last-val { font-size:14px; font-weight:800; }

.kpi-card { background:#FFFFFF; border-radius:16px; padding:16px 18px; box-shadow:0 2px 12px rgba(22,32,43,0.06); min-height:158px; }
.kpi-top { display:flex; align-items:center; justify-content:space-between; }
.kpi-label-top { font-size:11px; color:#6B7B8C; text-transform:uppercase; letter-spacing:.6px; font-weight:700; }
.kpi-num-big { font-size:34px; font-weight:800; margin-top:6px; line-height:1.1; }
.kpi-sub { font-size:12px; color:#6B7B8C; margin-top:6px; }
.chip { display:inline-block; padding:3px 10px; border-radius:20px; font-size:11px; font-weight:800; }
.panel-title { font-size:15px; font-weight:800; margin-bottom:2px; }
.panel-sub { font-size:12px; color:#6B7B8C; margin-bottom:8px; }

.leyenda-row { display:flex; align-items:center; justify-content:space-between; font-size:13px; padding:5px 0; }
.leyenda-row .dot { width:10px; height:10px; border-radius:50%; display:inline-block; margin-right:8px; }

.evento-row { display:flex; align-items:center; justify-content:space-between; padding:8px 12px; border-radius:8px;
    margin-bottom:6px; background:#F4F7FA; border-left:4px solid #ccc; font-family:Consolas,monospace; font-size:12px; }

.tabla-wrap { max-height:380px; overflow:auto; }
.tabla-alertas { width:100%; border-collapse:collapse; font-size:13px; }
.tabla-alertas th { text-align:left; font-size:11px; color:#6B7B8C; letter-spacing:.6px; padding:8px 10px; border-bottom:1px solid #EEF1F5; position:sticky; top:0; background:#fff; }
.tabla-alertas td { padding:11px 10px; border-bottom:1px solid #F1F4F8; }
.tabla-alertas .mono { font-family:Consolas,monospace; color:#6B7B8C; }

/* sidebar */
.sb-logo { display:flex; align-items:center; gap:10px; padding:4px 4px 18px 4px; }
.sb-logo .box { width:40px; height:40px; border-radius:11px; background:#EAF2FE; display:flex; align-items:center; justify-content:center; }
.sb-logo .box img { width:24px; }
.sb-logo b { font-size:15px; letter-spacing:.5px; display:block; line-height:1.1; }
.sb-logo small { font-size:9px; color:#8FA3B8 !important; letter-spacing:1px; }
.sb-user { display:flex; align-items:center; gap:10px; background:#1D2B3B; border-radius:12px; padding:10px 12px; margin:24px 0 12px 0; }
.sb-user .nm { font-size:13px; font-weight:700; line-height:1.2; }
.sb-user .rl { font-size:11px; color:#8FA3B8 !important; }
[data-testid="stSidebar"] .stButton button { background:transparent; border:1px solid #2E4058; color:#fff; border-radius:10px; }
[data-testid="stSidebar"] .stButton button:hover { background:#22324A; border-color:#2E7CE4; }
/* navegacion lateral con botones */
[class*="st-key-nav_"] button { background:transparent !important; border:none !important; justify-content:flex-start !important;
    padding:11px 14px !important; border-radius:11px !important; height:auto !important; min-height:0 !important;
    color:#FFFFFF !important; font-weight:600 !important; }
[class*="st-key-nav_"] button:hover { background:#22324A !important; }
[class*="st-key-nav_"] button[kind="primary"] { background:#2E7CE4 !important; font-weight:700 !important; }

/* detalle de incidente */
.info-grid { display:grid; grid-template-columns:1fr 1fr; gap:12px; margin-top:10px; }
.info-cell { background:#F6F8FB; border-radius:12px; padding:12px 14px; }
.info-val { font-size:15px; font-weight:800; margin-top:4px; }
.info-sub { font-size:11px; color:#6B7B8C; margin-top:3px; }
.step { display:flex; gap:12px; align-items:flex-start; padding:12px 14px; border:1px solid #EEF1F5; border-radius:12px; margin-bottom:10px; }
.step .n { width:26px; height:26px; border-radius:50%; background:#2E7CE4; color:#FFFFFF !important; display:flex;
    align-items:center; justify-content:center; font-weight:800; font-size:12px; flex-shrink:0; }
.step b { font-size:14px; }
.step p { margin:2px 0 0 0; font-size:12px; color:#6B7B8C; }
.aviso { background:#EEF4FF; border-radius:10px; padding:10px 12px; font-size:12px; color:#3B4A5C; margin-top:6px; }
.tl-row { display:flex; justify-content:space-between; gap:10px; padding:9px 0; border-bottom:1px solid #F1F4F8; font-size:13px; }
.tl-row small { color:#6B7B8C; white-space:nowrap; }
.tl-dot { width:9px; height:9px; border-radius:50%; display:inline-block; margin-right:8px; }
.term { background:#0F1722; border-radius:12px; padding:12px 16px; font-family:Consolas,monospace; font-size:12px; }
.term-row { display:flex; justify-content:space-between; padding:4px 0; border-bottom:1px solid #1B2A3D; }
.term-row span:first-child { color:#8FA3B8 !important; }
.term-row span:last-child { color:#CFE3FF !important; }
.nota { background:#F4F7FA; border-radius:12px; padding:10px 12px; margin-bottom:8px; font-size:13px; }
.nota small { color:#6B7B8C; }

/* reportes */
.tile-grid { display:grid; grid-template-columns:1fr 1fr; gap:12px; margin:12px 0; }
.tile { border-radius:14px; padding:14px 16px; }
.tile .lbl { font-size:11px; font-weight:800; letter-spacing:.6px; }
.tile .num { font-size:32px; font-weight:800; line-height:1.15; color:#16202B; }
.rep-head { display:flex; align-items:center; gap:12px; padding-bottom:12px; border-bottom:1px solid #EEF1F5; }
.rep-head img { width:34px; }
.rep-head b { font-size:16px; letter-spacing:.5px; }
.tag-proto { background:#E6EEFF; color:#1D4ED8 !important; font-size:10px; font-weight:800; padding:2px 8px; border-radius:6px; margin-left:6px; }
.st-key-dl_pdf button { background:#FDECEC !important; color:#C62828 !important; border:none !important; font-weight:700 !important; border-radius:10px !important; }
.st-key-dl_xlsx button { background:#E3F4E5 !important; color:#2E7D32 !important; border:none !important; font-weight:700 !important; border-radius:10px !important; }
@media print {
    [data-testid="stSidebar"], [data-testid="stHeader"], .topbar-card, .st-key-card_params, .st-key-acciones_rep { display:none !important; }
    .block-container { padding:0 !important; max-width:100% !important; }
}
/* texto de botones (el color real va en el <p> interno) */
.stButton button[kind="primary"] p, .stButton button[kind="primary"] span { color:#FFFFFF !important; font-weight:700 !important; }
.st-key-forgot button p, .st-key-ir_registro button p, .st-key-ir_login button p, .st-key-btn_ver_todas button p {
    color:#2E7CE4 !important; font-weight:700 !important; }
.st-key-forgot button, .st-key-ir_registro button, .st-key-ir_login button { justify-content:flex-end !important; }
.st-key-forgot button p, .st-key-ir_registro button p, .st-key-ir_login button p { text-align:right !important; }
.st-key-btn_ver_todas button { justify-content:flex-start !important; padding:8px 0 0 0 !important; }
.st-key-dl_pdf button p { color:#C62828 !important; font-weight:700 !important; }
.st-key-dl_xlsx button p { color:#2E7D32 !important; font-weight:700 !important; }
[class*="st-key-nav_"] button p { color:#FFFFFF !important; text-align:left !important; }
[class*="st-key-nav_"] button div { width:100% !important; text-align:left !important; }
[data-testid="stSidebar"] [data-testid="stVerticalBlock"] { gap:0.35rem; }
.kpi-card { min-height:180px; }
[class*="st-key-nav_"] button div, [class*="st-key-nav_"] button span,
.st-key-btn_ver_todas button div, .st-key-btn_ver_todas button span { justify-content:flex-start !important; }
.st-key-forgot button div, .st-key-forgot button span, .st-key-ir_registro button div, .st-key-ir_registro button span,
.st-key-ir_login button div, .st-key-ir_login button span { justify-content:flex-end !important; }
.stButton button p { margin:0 !important; }
.dot-live { width:8px; height:8px; border-radius:50%; background:#3C8C40; display:inline-block; margin-right:5px; animation:pulso 1.2s infinite; }
@keyframes pulso { 0% { opacity:1; } 50% { opacity:.2; } 100% { opacity:1; } }
.feed-row { display:flex; align-items:center; justify-content:space-between; padding:7px 12px; border-radius:8px; margin-bottom:5px;
    background:#F4F7FA; border-left:4px solid #ccc; font-family:Consolas,monospace; font-size:12px; }
.st-key-btn_sensor_toggle button[kind="secondary"] { border:1px solid #E23D3D; }
.st-key-btn_sensor_toggle button[kind="secondary"] p { color:#C62828 !important; font-weight:700 !important; }
.enlace-card { background:#FFFFFF; border-radius:16px; padding:16px 20px; margin-bottom:12px; box-shadow:0 2px 12px rgba(22,32,43,0.06); border-left:6px solid #ccc; }
.enlace-url { font-family:Consolas,monospace; font-size:13px; background:#F4F7FA; border-radius:8px; padding:8px 10px; margin:10px 0; word-break:break-all; }
.barra { height:8px; border-radius:6px; background:#EEF1F5; overflow:hidden; margin:6px 0 4px 0; }
.barra div { height:100%; border-radius:6px; }
.razon { font-size:13px; padding:4px 0; display:flex; gap:8px; align-items:flex-start; }
.razon .pt { width:9px; height:9px; border-radius:50%; margin-top:5px; flex-shrink:0; }
.que-hacer { background:#EEF4FF; border-radius:10px; padding:10px 12px; font-size:13px; margin-top:10px; }
.area { font-size:10px; font-weight:800; color:#6B7B8C !important; text-transform:uppercase; letter-spacing:.5px; min-width:92px; padding-top:2px; }
.aviso-amarillo { background:#FFF8E1; border-radius:10px; padding:10px 12px; font-size:12px; color:#6B4E00; margin-top:8px; }
</style>
""".replace("__FONDO__", fondo_css)
st.markdown(CSS, unsafe_allow_html=True)


def encabezado_tarjeta(titulo, desc):
    st.markdown(f'<div class="logo-box"><img src="data:image/png;base64,{logo_b64}"></div>', unsafe_allow_html=True)
    st.markdown(f'<div class="card-title">{titulo}</div>', unsafe_allow_html=True)
    st.markdown('<div class="card-subtitle">Detección de incidentes para pymes</div>', unsafe_allow_html=True)
    st.markdown(f'<div class="card-desc">{desc}</div>', unsafe_allow_html=True)


def pie_tarjeta():
    st.markdown('<div class="footer-sub">Protegido por <b>motor de detección con Machine Learning</b></div>', unsafe_allow_html=True)
    st.markdown('<div class="footer-mono">Prototipo académico · RF &amp; XGBoost (NSL-KDD)</div>', unsafe_allow_html=True)


def ir_a(pagina_destino):
    st.session_state["pagina"] = pagina_destino


def ir_seccion(s):
    st.session_state["seccion"] = s
    st.session_state.pop("incidente_abierto", None)


def liberar_incidente(usuario, inc_id, mensaje_id, autor):
    cfg = st.session_state.get("buzon_cfg")
    if not cfg or imap_conector is None:
        st.session_state["toast"] = "Primero conecta el buzón en «Conector de correo» para poder liberar el correo."
        return
    try:
        with imap_conector.BuzonIMAP(cfg["usuario"], cfg["clave"]) as buzon:
            buzon.liberar(mensaje_id)
    except imap_conector.ErrorBuzon as e:
        st.session_state["toast"] = "No se pudo liberar el correo: " + str(e)[:140]
        return
    cambiar_estado_incidente(usuario, inc_id, "Liberado", autor)
    st.session_state["toast"] = "Listo: el correo volvió a la bandeja de entrada."


def abrir_incidente(inc_id):
    st.session_state["incidente_abierto"] = inc_id


def cerrar_incidente():
    st.session_state.pop("incidente_abierto", None)


def publicar_nota(usuario, inc_id, autor):
    texto = st.session_state.get("nota_txt", "").strip()
    if texto:
        agregar_nota(usuario, inc_id, autor, texto[:500])
    st.session_state["nota_txt"] = ""


def mostrar_html(contenido, alto=44):
    if hasattr(st, "iframe"):
        st.iframe(contenido, height=alto)
    else:
        components.html(contenido, height=alto)


def chip(texto, estilo, punto=False):
    fg, bg = estilo
    return ('<span class="chip" style="color:' + fg + ';background:' + bg + ';">'
            + ("● " if punto else "") + html.escape(texto) + '</span>')


def figura_donut(dist, total, alto=190):
    fig = go.Figure(go.Pie(labels=list(dist.keys()), values=list(dist.values()), hole=0.68,
                           marker=dict(colors=[COLOR_CAT[c] for c in dist]), textinfo="none", sort=False))
    fig.add_annotation(text="<b>" + str(total) + "</b><br><span style='font-size:10px'>eventos</span>",
                       showarrow=False, font=dict(size=18))
    fig.update_layout(height=alto, margin=dict(l=0, r=0, t=0, b=0), paper_bgcolor="rgba(0,0,0,0)", showlegend=False)
    return fig


def leyenda_html(dist, total):
    filas = ""
    for c, n in dist.items():
        pct = round(n / total * 100) if total else 0
        nombre = "Sin amenaza" if c == "Normal" else c
        filas += ('<div class="leyenda-row"><span><span class="dot" style="background:' + COLOR_CAT[c] + '"></span>'
                  + nombre + '</span><b>' + str(pct) + '%</b></div>')
    return filas


def vista_historial(usuario):
    st.markdown('<div class="dash-title">↻ Historial de incidentes</div>'
                '<div class="dash-sub">Todos los incidentes detectados. Marca la casilla de una fila para ver su detalle y cambiar su estado.</div><br>',
                unsafe_allow_html=True)
    with st.container(key="card_filtros"):
        f1, f2, f3, f4 = st.columns([1.3, 1.3, 1, 1])
        with f1:
            cats = st.multiselect("Categoría", CATEGORIAS_TODAS, default=CATEGORIAS_TODAS, key="h_cats")
        with f2:
            estados = st.multiselect("Estado", ESTADOS, default=[], placeholder="Todos", key="h_estados")
        with f3:
            desde = st.date_input("Desde", value=date.today() - timedelta(days=30), format="DD/MM/YYYY", key="h_desde")
        with f4:
            hasta = st.date_input("Hasta", value=date.today(), format="DD/MM/YYYY", key="h_hasta")

    incs = listar_incidentes(usuario, desde=desde, hasta=hasta, cats=cats, estados=estados or None)
    with st.container(key="card_hist_tabla"):
        if not incs:
            st.info("No hay incidentes con esos filtros. Sube un archivo en el Dashboard para generar un análisis.")
            return
        st.markdown('<div class="panel-title">' + str(len(incs)) + ' incidente(s)</div>', unsafe_allow_html=True)
        df = pd.DataFrame([{
            "ID": i["id"], "Fecha": fmt_fecha(i["fecha"]), "Tipo de amenaza": NOMBRES_AMENAZA[i["categoria"]],
            "Activo (simulado)": i["activo"], "Severidad": i["severidad"], "Estado": i["estado"],
            "Certeza": "{:.0f}%".format(i["certeza"] * 100),
            "Origen": {"sensor": "Sensor virtual", "correo": "Correo subido", "buzon": "Buzón de correo"}.get(i.get("origen"), "Archivo CSV"),
        } for i in incs])

        def color_sev(v):
            fg, bg = SEV_ESTILO.get(v, ("#16202B", "#FFFFFF"))
            return "color:" + fg + ";background-color:" + bg + ";font-weight:700"

        def color_est(v):
            fg, bg = ESTADO_ESTILO.get(v, ("#16202B", "#FFFFFF"))
            return "color:" + fg + ";background-color:" + bg + ";font-weight:700"

        estilo = df.style.map(color_sev, subset=["Severidad"]).map(color_est, subset=["Estado"])
        ev = st.dataframe(estilo, hide_index=True, width="stretch", on_select="rerun",
                          selection_mode="single-row", key="tabla_hist")
        if ev.selection.rows:
            st.session_state["incidente_abierto"] = int(df.iloc[ev.selection.rows[0]]["ID"])
            st.rerun()


def _secretos():
    try:
        return st.secrets
    except Exception:
        return None


def _contexto_chat():
    """Si hay un incidente abierto en pantalla, el chat habla de ese incidente."""
    usuario = st.session_state.get("usuario")
    inc_id = st.session_state.get("incidente_abierto")
    if st.session_state.get("pagina") == "dashboard" and usuario and st.session_state.get("seccion") == "historial" and inc_id:
        inc = obtener_incidente(usuario, inc_id)
        if inc is not None:
            try:
                det = json.loads(inc["detalle"] or "{}")
            except Exception:
                det = {}
            return dict(inc), det
    return None, None


def _historial_chat(inc):
    return st.session_state.setdefault("chat_inc_" + str(inc["id"]) if inc else "chat_ayuda", [])


def enviar_chat(pregunta=None):
    if pregunta is None:
        pregunta = st.session_state.get("chat_flot_txt", "").strip()
        st.session_state["chat_flot_txt"] = ""
    if not pregunta or asistente is None or st.session_state.get("chat_pendiente"):
        return
    inc, _ = _contexto_chat()
    _historial_chat(inc).append({"rol": "user", "texto": pregunta})
    st.session_state["chat_pendiente"] = pregunta


PUNTITOS_HTML = '<div class="typing"><span></span><span></span><span></span></div>'


def _respuesta_local(pregunta, inc, det, previo, logueado):
    if inc:
        return asistente.responder(pregunta, inc, det, NOMBRES_AMENAZA[inc["categoria"]], ACCIONES_CAT[inc["categoria"]], previo, usar_ia=False)[0]
    return asistente.ayuda_sin_internet(pregunta, logueado)


def _responder_pendiente(inc, det, historial):
    """Muestra los puntitos y luego el texto apareciendo poco a poco."""
    pregunta = st.session_state.get("chat_pendiente")
    previo = historial[:-1]
    usadas = st.session_state.get("chat_ia_usadas", 0)
    pagina_act = st.session_state.get("pagina", "login")
    logueado = pagina_act == "dashboard"
    pantalla = {"login": "inicio de sesión", "registro": "registro de empresa"}.get(pagina_act, "sección " + str(st.session_state.get("seccion", "dashboard")))
    caja = st.empty()
    caja.markdown(PUNTITOS_HTML, unsafe_allow_html=True)
    texto, modo, aviso = "", "local", None
    flujo = None
    if usadas < asistente.MAX_PREGUNTAS_IA and st.session_state.get("pref_ia", True):
        try:
            flujo = (asistente.flujo_incidente(pregunta, inc, det, NOMBRES_AMENAZA[inc["categoria"]], previo, secrets=_secretos()) if inc
                     else asistente.flujo_ayuda(pregunta, previo, pantalla, logueado, secrets=_secretos()))
        except Exception:
            flujo = None
    elif asistente.clave_gemini(_secretos()):
        aviso = "Se alcanzó el máximo de preguntas con IA en esta sesión."
    if flujo is not None:
        try:
            for trozo in flujo:
                texto += trozo
                caja.markdown(texto + " ▌")
            modo = "gemini"
            st.session_state["chat_ia_usadas"] = usadas + 1
        except asistente.ErrorGemini as e:
            aviso = "Respondí en modo sin internet porque " + str(e) + "."
        except Exception:
            aviso = "Respondí en modo sin internet porque hubo un problema con Gemini."
    if not texto:
        modo = "local"
        try:
            completo = _respuesta_local(pregunta, inc, det, previo, logueado)
        except Exception:
            completo = "Tuve un problema para responder. Intenta de nuevo en un momento."
        time.sleep(0.5)
        palabras = completo.split(" ")
        for k in range(0, len(palabras), 3):
            caja.markdown(" ".join(palabras[:k + 3]) + " ▌")
            time.sleep(0.03)
        texto = completo
    caja.markdown(texto)
    historial.append({"rol": "bot", "texto": texto, "modo": modo, "aviso": aviso})
    st.session_state.pop("chat_pendiente", None)


def alternar_chat():
    st.session_state["chat_abierto"] = not st.session_state.get("chat_abierto", False)


def borrar_chat_flotante():
    inc, _ = _contexto_chat()
    st.session_state.pop("chat_inc_" + str(inc["id"]) if inc else "chat_ayuda", None)


CSS_CHAT = """
<style>
.st-key-chat_root { position: fixed; top: 0; left: 0; width: 0 !important; height: 0; overflow: visible; z-index: 10000; }
.st-key-chat_panel input { padding-left: 12px !important; background-image: none !important; }
.st-key-chat_fab_wrap { position: fixed; bottom: 22px; right: 22px; z-index: 10001; width: auto !important; }
.st-key-chat_fab_wrap button { border-radius: 999px !important; border: none; padding: 12px 22px; font-weight: 700; color: #FFFFFF;
    background: linear-gradient(135deg, #2E7CE4 0%, #7B61FF 100%); box-shadow: 0 8px 24px rgba(46,124,228,.40); }
.st-key-chat_fab_wrap button:hover { color: #FFFFFF; filter: brightness(1.08); border: none; }
.st-key-chat_fab_wrap button p { color: #FFFFFF; font-weight: 700; }
.st-key-chat_panel { position: fixed; bottom: 86px; right: 22px; z-index: 10000; width: min(390px, calc(100vw - 32px)) !important;
    max-height: calc(100vh - 110px); overflow-y: auto; background: #FFFFFF; border: 1px solid #E3E8EF; border-radius: 18px;
    box-shadow: 0 14px 44px rgba(15,30,60,.28); padding: 0 14px 12px 14px; gap: 8px; }
.chat-head { margin: 0 -14px 4px -14px; padding: 16px 18px; color: #FFFFFF; border-radius: 18px 18px 0 0;
    background: linear-gradient(135deg, #2E7CE4 0%, #7B61FF 100%); display: flex; gap: 12px; align-items: center; }
.chat-head img { width: 38px; height: 38px; background: #FFFFFF; border-radius: 10px; padding: 4px; }
.chat-head b { font-size: 17px; display: block; }
.chat-head small { font-size: 12px; opacity: .92; line-height: 1.35; display: block; }
.chat-head .tag { font-size: 10px; font-weight: 700; background: rgba(255,255,255,.22); border-radius: 6px; padding: 1px 6px; margin-left: 6px; vertical-align: middle; }
.chat-ctx { font-size: 12px; background: #F1ECFF; color: #4C35C9; border-radius: 10px; padding: 6px 10px; }
.st-key-chat_panel [class*="st-key-chat_op_"] button { justify-content: flex-start; text-align: left; border: 1px solid #B9A8FF;
    background: #F7F4FF; border-radius: 12px; color: #2B2350; }
.st-key-chat_panel [class*="st-key-chat_op_"] button:hover { background: #EDE6FF; border-color: #7B61FF; color: #2B2350; }
.st-key-chat_panel [data-testid="stChatMessage"] { padding: 6px 8px; font-size: 14px; }
.typing { display: inline-flex; gap: 5px; padding: 10px 14px; background: #F1ECFF; border-radius: 14px; }
.typing span { width: 8px; height: 8px; border-radius: 50%; background: #7B61FF; display: inline-block; animation: chatpunto 1.2s infinite ease-in-out; }
.typing span:nth-child(2) { animation-delay: .2s; }
.typing span:nth-child(3) { animation-delay: .4s; }
@keyframes chatpunto { 0%, 80%, 100% { opacity: .25; transform: translateY(0); } 40% { opacity: 1; transform: translateY(-4px); } }
</style>
"""


def chat_flotante():
    if asistente is None:
        return
    with st.container(key="chat_root"):
        _chat_flotante_contenido()


@st.fragment
def _chat_flotante_contenido():
    abierto = st.session_state.get("chat_abierto", False)
    st.markdown(CSS_CHAT, unsafe_allow_html=True)
    with st.container(key="chat_fab_wrap"):
        st.button("✕  Cerrar" if abierto else "💬  Chatea con Centinela", key="chat_fab", on_click=alternar_chat)
    if not abierto:
        return
    inc, det = _contexto_chat()
    logueado = st.session_state.get("pagina") == "dashboard"
    hay_ia = bool(asistente.clave_gemini(_secretos()))
    historial = _historial_chat(inc)
    with st.container(key="chat_panel"):
        st.markdown('<div class="chat-head"><img src="data:image/png;base64,' + logo_b64 + '"><div><b>Chatea con Centinela'
                    '<span class="tag">' + ("IA" if hay_ia else "SIN INTERNET") + '</span></b>'
                    '<small>Te ayudo a usar la plataforma y a entender las alertas.</small></div></div>', unsafe_allow_html=True)
        if inc:
            st.markdown('<div class="chat-ctx">Hablando del <b>incidente #' + str(inc["id"]) + '</b> · ' + html.escape(NOMBRES_AMENAZA[inc["categoria"]])
                        + ' · riesgo ' + html.escape(inc["severidad"].lower()) + '</div>', unsafe_allow_html=True)
        with st.container(height=340, key="chat_msgs", border=False):
            with st.chat_message("assistant", avatar="🛡️"):
                if inc:
                    st.markdown("¡Hola! 👋 Soy **Centinela**. Pregúntame lo que quieras sobre este incidente.")
                elif logueado:
                    st.markdown("¡Hola! 👋 Soy **Centinela**. ¿En qué te ayudo? Elige una opción o escríbeme.")
                else:
                    st.markdown("¡Hola! 👋 Soy **Centinela**. ¿Tienes dudas para entrar o registrarte? Elige una opción o escríbeme.")
            for m in historial[-12:]:
                with st.chat_message("user" if m["rol"] == "user" else "assistant", avatar="🧑" if m["rol"] == "user" else "🛡️"):
                    st.markdown(m["texto"])
                    if m["rol"] == "bot":
                        st.caption(("Con IA (Gemini)" if m.get("modo") == "gemini" else "Modo sin internet")
                                   + (" · " + m["aviso"] if m.get("aviso") else ""))
            if st.session_state.get("chat_pendiente"):
                with st.chat_message("assistant", avatar="🛡️"):
                    _responder_pendiente(inc, det, historial)
                st.rerun(scope="fragment")
            if not historial:
                opciones = asistente.PREGUNTAS_RAPIDAS if inc else (asistente.PREGUNTAS_AYUDA_APP if logueado else asistente.PREGUNTAS_AYUDA_LOGIN)
                for i, p in enumerate(opciones):
                    st.button(p, key="chat_op_" + str(i), on_click=enviar_chat, args=(p,), width="stretch")
        try:
            components.html("<script>(function(){var n=" + str(len(historial)) + ";function bajar(){var d=window.parent.document;"
                            "var el=d.querySelector('.st-key-chat_msgs');if(!el)return;var x=el;for(var k=0;k<6&&x&&!(x.classList&&x.classList.contains('st-key-chat_root'));k++){x.scrollTop=x.scrollHeight;x=x.parentElement;}"
                            "el.querySelectorAll('*').forEach(function(c){if(c.scrollHeight>c.clientHeight+5){c.scrollTop=c.scrollHeight;}});}"
                            "setTimeout(bajar,80);setTimeout(bajar,400);})();</script>", height=0)
        except Exception:
            pass
        c1, c2 = st.columns([5, 1], vertical_alignment="center")
        c1.text_input("Mensaje", key="chat_flot_txt", label_visibility="collapsed", placeholder="Hacer una pregunta...", on_change=enviar_chat)
        c2.button("➤", key="chat_enviar", type="primary", on_click=enviar_chat, width="stretch")
        p1, p2 = st.columns([3, 1.3], vertical_alignment="center")
        p1.caption("No escribas contraseñas ni datos personales." + (" Solo se envían datos mínimos a Gemini." if hay_ia else ""))
        p2.button("Borrar chat", key="chat_borrar", type="tertiary", on_click=borrar_chat_flotante, disabled=not historial)


def accion_guardar_empresa(usuario):
    ok, msg = cambiar_empresa(usuario["id"], st.session_state.get("cfg_empresa", ""))
    if ok:
        st.session_state["usuario"]["empresa"] = st.session_state["cfg_empresa"].strip()[:120]
    st.session_state["cfg_msg_empresa"] = (ok, msg)


def accion_cambiar_clave(usuario):
    ok, msg = cambiar_clave(usuario["id"], st.session_state.get("cfg_clave_act", ""),
                            st.session_state.get("cfg_clave_new", ""), st.session_state.get("cfg_clave_new2", ""))
    st.session_state["cfg_msg_clave"] = (ok, msg)
    if ok:
        for k in ("cfg_clave_act", "cfg_clave_new", "cfg_clave_new2"):
            st.session_state[k] = ""


def borrar_conversaciones():
    for k in [k for k in st.session_state.keys() if str(k).startswith("chat_inc_") or k == "chat_ayuda"]:
        st.session_state.pop(k, None)
    st.session_state["toast"] = "Se borraron las conversaciones con el asistente."


def _mensaje(clave):
    r = st.session_state.pop(clave, None)
    if r:
        (st.success if r[0] else st.error)(r[1])


def vista_configuracion(usuario):
    es_admin = usuario["rol"] == "admin"
    cuenta, n_inc, n_corr = datos_cuenta(usuario["id"])
    st.markdown('<div class="dash-title">⚙ Configuración</div>'
                '<div class="dash-sub">Tu cuenta, tu contraseña, el asistente y la privacidad de tus datos.</div><br>', unsafe_allow_html=True)
    izq, der = st.columns([1.15, 1])
    with izq:
        with st.container(key="card_cfg_cuenta"):
            celdas = [("EMPRESA", html.escape(cuenta.get("empresa", "")), "Así aparece en reportes y en la barra superior"),
                      ("CORREO DE ACCESO", html.escape(cuenta.get("correo", "")), "Con este correo inicias sesión"),
                      ("TIPO DE CUENTA", "Administrador" if es_admin else "Usuario TI", "Estado: " + html.escape(cuenta.get("estado", ""))),
                      ("REGISTRADA DESDE", html.escape(str(cuenta.get("fecha", ""))[:10]),
                       str(n_inc) + " incidentes · " + str(n_corr) + " correos revisados")]
            grid = "".join('<div class="info-cell"><div class="kpi-label-top">' + t + '</div><div class="info-val">' + v
                           + '</div><div class="info-sub">' + s + '</div></div>' for t, v, s in celdas)
            st.markdown('<div class="panel-title">Datos de la cuenta</div><div class="info-grid">' + grid + '</div>', unsafe_allow_html=True)
            if not es_admin:
                st.write("")
                st.markdown('<div class="panel-sub"><b>Cambiar el nombre de la empresa</b></div>', unsafe_allow_html=True)
                if "cfg_empresa" not in st.session_state:
                    st.session_state["cfg_empresa"] = cuenta.get("empresa", "")
                e1, e2 = st.columns([3, 1], vertical_alignment="bottom")
                e1.text_input("Empresa", key="cfg_empresa", label_visibility="collapsed")
                e2.button("Guardar", key="cfg_btn_empresa", on_click=accion_guardar_empresa, args=(usuario,), width="stretch")
                _mensaje("cfg_msg_empresa")
        with st.container(key="card_cfg_clave"):
            st.markdown('<div class="panel-title">🔒 Cambiar contraseña</div>'
                        '<div class="panel-sub">Mínimo 8 caracteres. Usa una que no uses en otros sitios.</div>', unsafe_allow_html=True)
            if es_admin:
                st.info("La contraseña del administrador se define en la configuración del servidor, no desde aquí.")
            else:
                st.text_input("Contraseña actual", key="cfg_clave_act", type="password", placeholder="Contraseña actual", label_visibility="collapsed")
                st.text_input("Nueva contraseña", key="cfg_clave_new", type="password", placeholder="Nueva contraseña", label_visibility="collapsed")
                st.text_input("Repite la nueva contraseña", key="cfg_clave_new2", type="password", placeholder="Repite la nueva contraseña", label_visibility="collapsed")
                st.button("Cambiar contraseña", key="cfg_btn_clave", type="primary", on_click=accion_cambiar_clave, args=(usuario,))
                _mensaje("cfg_msg_clave")
    with der:
        with st.container(key="card_cfg_asistente"):
            hay_ia = asistente is not None and bool(asistente.clave_gemini(_secretos()))
            st.markdown('<div class="panel-title">💬 Asistente «Chatea con Centinela»</div>'
                        '<div class="panel-sub">' + ("IA en línea configurada (Gemini). Si no hay internet, responde en modo sin internet."
                                                     if hay_ia else "Sin clave de IA: responde en modo sin internet (nada sale de este equipo).") + '</div>',
                        unsafe_allow_html=True)
            if hay_ia:
                st.toggle("Usar IA en línea (Gemini)", key="pref_ia", value=st.session_state.get("pref_ia", True),
                          help="Si la apagas, el asistente responde solo con información guardada en Centinela, sin enviar nada a internet.")
                usadas = st.session_state.get("chat_ia_usadas", 0)
                st.caption("Preguntas con IA en esta sesión: " + str(usadas) + " de " + str(asistente.MAX_PREGUNTAS_IA) + ".")
            st.button("Borrar mis conversaciones con el asistente", key="cfg_btn_chat", on_click=borrar_conversaciones)
        with st.container(key="card_cfg_privacidad"):
            st.markdown('<div class="panel-title">🛡️ Privacidad de tus datos</div>'
                        '<ul style="font-size:13px;line-height:1.65;padding-left:18px;margin:6px 0 0 0;">'
                        '<li>De cada correo se guarda <b>solo</b> el remitente, el asunto, las razones y los puntajes. <b>Nunca el contenido.</b></li>'
                        '<li>La contraseña de aplicación de Gmail <b>no se guarda</b>: vive solo mientras la sesión está abierta.</li>'
                        '<li>Las contraseñas de las cuentas se guardan <b>cifradas</b>.</li>'
                        '<li>Cada empresa ve <b>solo sus propios</b> incidentes.</li>'
                        '<li>Al asistente con IA se le envían datos mínimos, sin correos, números ni el texto de los mensajes.</li>'
                        '<li>Para quitarle el acceso a tu Gmail, borra la contraseña de aplicación en myaccount.google.com/apppasswords.</li>'
                        '</ul>', unsafe_allow_html=True)
        with st.container(key="card_cfg_acerca"):
            st.markdown('<div class="panel-title">ℹ️ Acerca de Centinela</div>'
                        '<div class="panel-sub" style="line-height:1.6;">Versión 1.0 · Octubre de 2026<br>'
                        'Detección de phishing en correos (texto en español entrenado con SpaPhish, enlaces y señales SPF, DKIM y DMARC), '
                        'verificador de enlaces, conector de Gmail con cuarentena y avisos, y módulo de tráfico de red (NSL-KDD).<br>'
                        'Trabajo de grado · Ingeniería de Sistemas · Universidad Católica Luis Amigó.</div>', unsafe_allow_html=True)


def vista_incidente(usuario, inc_id, autor):
    inc = obtener_incidente(usuario, inc_id)
    if inc is None:
        st.session_state.pop("incidente_abierto", None)
        st.error("No se encontró ese incidente.")
        return
    cat = inc["categoria"]
    st.markdown('<div class="dash-sub">Dashboard › Historial › <b>Incidente #' + str(inc["id"]) + '</b></div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="dash-title" style="margin:6px 0 8px 0;">Incidente #' + str(inc["id"]) + ' · ' + NOMBRES_AMENAZA[cat] + '</div>'
        '<div style="margin-bottom:8px;">' + chip(inc["severidad"], SEV_ESTILO[inc["severidad"]], True) + ' &nbsp;'
        + chip(inc["estado"], ESTADO_ESTILO[inc["estado"]]) + '</div>'
        '<div class="dash-sub" style="max-width:820px;">' + DESCRIPCION_CAT[cat] + '</div>', unsafe_allow_html=True)
    st.write("")

    b = st.columns([1, 1.4, 1.8, 1, 1.7])
    b[0].button("← Volver", key="inc_volver", on_click=cerrar_incidente, width="stretch")
    b[1].button("Tomar en revisión", key="inc_rev", on_click=cambiar_estado_incidente,
                args=(usuario, inc_id, "En revisión", autor), disabled=inc["estado"] == "En revisión", width="stretch")
    b[2].button("✔ Marcar como resuelto", key="inc_res", type="primary", on_click=cambiar_estado_incidente,
                args=(usuario, inc_id, "Resuelto", autor), disabled=inc["estado"] == "Resuelto", width="stretch")
    b[3].button("Escalar", key="inc_esc", on_click=cambiar_estado_incidente,
                args=(usuario, inc_id, "Escalado", autor), disabled=inc["estado"] == "Escalado", width="stretch")
    b[4].button("Descartar falso positivo", key="inc_fp", on_click=cambiar_estado_incidente,
                args=(usuario, inc_id, "Falso positivo", autor), disabled=inc["estado"] == "Falso positivo", width="stretch")
    st.write("")

    try:
        detalle = json.loads(inc["detalle"] or "{}")
    except Exception:
        detalle = {}
    if cat == "Phishing" and detalle.get("mensaje_id"):
        st.button("↩ Liberar el correo a la bandeja de entrada", key="inc_liberar", on_click=liberar_incidente,
                  args=(usuario, inc_id, detalle["mensaje_id"], autor), disabled=inc["estado"] != "En cuarentena")
        st.write("")
    hist = historial_incidente(inc_id)
    izq, der = st.columns([1.6, 1])

    with izq:
        with st.container(key="card_inc_info"):
            celdas = [
                ("ORIGEN (IP)", "No disponible", "El dataset NSL-KDD no incluye direcciones IP"),
                ("ACTIVO AFECTADO", html.escape(inc["activo"]), "Simulado para la demostración"),
                ("MOMENTO DE DETECCIÓN", fmt_fecha(inc["fecha"]), "Evento #" + str(inc["evento_num"]) + " del archivo analizado"),
                ("MOTOR ANALÍTICO", "Modelo ML (XGBoost)", "Certeza del modelo: {:.0f}%".format(inc["certeza"] * 100)),
            ]
            if cat == "Phishing":
                celdas = [
                    ("REMITENTE", html.escape(str(detalle.get("remitente", "No disponible"))[:60]), "Dirección que figura en el correo (puede estar falsificada)"),
                    ("BUZÓN ANALIZADO", html.escape(inc["activo"]), "Correo recibido por la empresa"),
                    ("MOMENTO DE DETECCIÓN", fmt_fecha(inc["fecha"]), "Asunto: " + html.escape(str(detalle.get("asunto", ""))[:48])),
                    ("MOTOR ANALÍTICO", "Texto + enlaces + señales", "Riesgo estimado: {:.0f}%".format(inc["certeza"] * 100)),
                ]
            grid = "".join('<div class="info-cell"><div class="kpi-label-top">' + t + '</div><div class="info-val">' + v
                           + '</div><div class="info-sub">' + s + '</div></div>' for t, v, s in celdas)
            st.markdown('<div class="panel-title">Información del evento y detección</div><div class="info-grid">' + grid + '</div>',
                        unsafe_allow_html=True)
        with st.container(key="card_inc_accion"):
            pasos = "".join('<div class="step"><div class="n">' + str(n) + '</div><div><b>' + t + '</b><p>' + d + '</p></div></div>'
                            for n, (t, d) in enumerate(ACCIONES_CAT[cat], start=1))
            st.markdown('<div class="panel-title">Acción recomendada para contener la amenaza</div>'
                        '<div class="panel-sub">Guía general para el equipo de TI según el tipo de amenaza.</div>' + pasos +
                        '<div class="aviso">Son recomendaciones generales. CENTINELA <b>no ejecuta cambios en la red</b>: detecta, prioriza y sugiere.</div>',
                        unsafe_allow_html=True)
        with st.container(key="card_inc_tecnico"):
            filas = "".join('<div class="term-row"><span>' + html.escape(str(k)) + '</span><span>' + html.escape(str(v)) + '</span></div>'
                            for k, v in detalle.items())
            st.markdown('<div class="panel-title">Registro técnico del evento</div>'
                        '<div class="panel-sub">' + ("Datos que Centinela extrajo del correo. No se guarda el contenido del mensaje." if cat == "Phishing" else "Variables del evento que recibió el modelo (NSL-KDD).") + '</div>'
                        '<div class="term">' + (filas or "Sin datos") + '</div>', unsafe_allow_html=True)

    with der:
        with st.container(key="card_inc_linea"):
            filas = ""
            for h in [x for x in hist if x["tipo"] in ("sistema", "estado")]:
                color = "#E23D3D" if h["tipo"] == "sistema" else "#2E7CE4"
                filas += ('<div class="tl-row"><span><span class="tl-dot" style="background:' + color + '"></span>'
                          + html.escape(h["texto"]) + '</span><small>' + fmt_fecha(h["fecha"], corto=True) + '</small></div>')
            st.markdown('<div class="panel-title">Línea de tiempo</div>' + filas, unsafe_allow_html=True)
        with st.container(key="card_inc_notas"):
            notas = [x for x in hist if x["tipo"] == "nota"]
            cuerpo = "".join('<div class="nota"><b>' + html.escape(n["autor"]) + '</b> <small>' + fmt_fecha(n["fecha"])
                             + '</small><br>' + html.escape(n["texto"]) + '</div>' for n in notas)
            st.markdown('<div class="panel-title">Notas del equipo TI</div>' + (cuerpo or '<div class="panel-sub">Aún no hay notas.</div>'),
                        unsafe_allow_html=True)
            st.text_area("Nota", key="nota_txt", placeholder="Escribe hallazgos o acuerdos del equipo...",
                         label_visibility="collapsed", height=90)
            st.button("Publicar nota", key="btn_nota", type="primary", on_click=publicar_nota,
                      args=(usuario, inc_id, autor), width="stretch")
        with st.container(key="card_inc_integridad"):
            st.markdown('<div class="panel-title">🔒 Integridad del historial</div>'
                        '<div class="panel-sub">Este incidente queda registrado y no puede editarse: únicamente su estado se actualiza '
                        '(en revisión, resuelto, escalado o falso positivo) y cada cambio queda en la línea de tiempo.</div>',
                        unsafe_allow_html=True)


def vista_reportes(usuario, empresa, autor):
    st.markdown('<div class="dash-title">▤ Reportes</div>'
                '<div class="dash-sub">Genera un resumen de los incidentes detectados en un periodo.</div><br>', unsafe_allow_html=True)
    with st.container(key="card_params"):
        st.markdown('<div class="panel-title">Parámetros</div>', unsafe_allow_html=True)
        p1, p2 = st.columns(2)
        with p1:
            desde = st.date_input("Desde", value=date.today() - timedelta(days=30), format="DD/MM/YYYY", key="rep_desde")
        with p2:
            hasta = st.date_input("Hasta", value=date.today(), format="DD/MM/YYYY", key="rep_hasta")
        cc = st.columns(5)
        cats = [c for c, col in zip(CATEGORIAS_TODAS, cc) if col.checkbox(c, value=True, key="rep_cat_" + c)]
        generar = st.button("Generar reporte", key="btn_generar", type="primary", width="stretch")

    if generar:
        if desde > hasta:
            st.error("La fecha «Desde» no puede ser posterior a «Hasta».")
        else:
            st.session_state["reporte"] = {"desde": desde, "hasta": hasta, "cats": cats}
    cfg = st.session_state.get("reporte")
    if not cfg:
        return

    rep = calcular_reporte(usuario, cfg["desde"], cfg["hasta"], cfg["cats"])
    periodo = cfg["desde"].strftime("%d/%m/%Y") + " al " + cfg["hasta"].strftime("%d/%m/%Y")
    nombre_arch = "reporte_centinela_" + cfg["desde"].strftime("%Y%m%d") + "_" + cfg["hasta"].strftime("%Y%m%d")

    with st.container(key="acciones_rep"):
        a, b, c, d = st.columns([2.6, 1, 1, 1])
        with a:
            st.markdown('<div class="panel-title" style="margin-top:8px;">Reporte ' + periodo + '</div>', unsafe_allow_html=True)
        try:
            pdf = generar_pdf(empresa, autor, cfg["desde"], cfg["hasta"], rep)
            xlsx = generar_excel(empresa, autor, cfg["desde"], cfg["hasta"], rep)
            b.download_button("PDF", pdf, file_name=nombre_arch + ".pdf", mime="application/pdf", key="dl_pdf", width="stretch")
            c.download_button("Excel", xlsx, file_name=nombre_arch + ".xlsx",
                              mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="dl_xlsx", width="stretch")
        except ImportError:
            st.warning("Para descargar PDF y Excel instala en la terminal: pip install reportlab openpyxl")
        with d:
            mostrar_html(
                '<button onclick="window.parent.print()" style="width:100%;height:38px;border:1px solid #CBD5E1;border-radius:10px;'
                'background:#fff;font-weight:700;cursor:pointer;font-family:sans-serif;color:#16202B;">Imprimir</button>', alto=44)

    t = rep["tiles"]
    with st.container(key="card_reporte"):
        st.markdown(
            '<div class="rep-head"><img src="data:image/png;base64,' + logo_b64 + '"><div><b>CENTINELA</b>'
            '<span class="tag-proto">Prototipo</span><div class="dash-sub">' + html.escape(empresa) + ' · ' + periodo + '</div></div></div>'
            '<div class="tile-grid">'
            '<div class="tile" style="background:#FDECEC;"><div class="lbl" style="color:#C62828;">CRÍTICOS</div><div class="num">' + str(t["criticos"]) + '</div></div>'
            '<div class="tile" style="background:#FFF1E0;"><div class="lbl" style="color:#C2410C;">ALTOS</div><div class="num">' + str(t["altos"]) + '</div></div>'
            '<div class="tile" style="background:#FFF8E1;"><div class="lbl" style="color:#946200;">MEDIOS</div><div class="num">' + str(t["medios"]) + '</div></div>'
            '<div class="tile" style="background:#E8EDF3;"><div class="lbl" style="color:#475569;">FALSOS POSITIVOS</div><div class="num">' + str(t["fp"]) + '</div></div>'
            '</div><div class="panel-title">Distribución por categoría</div>', unsafe_allow_html=True)
        if rep["total"] == 0:
            st.info("No hay análisis en ese periodo. Sube un archivo en el Dashboard para generar datos.")
        else:
            d1, d2 = st.columns([1, 1.3])
            with d1:
                st.plotly_chart(figura_donut(rep["dist"], rep["total"], 200), width="stretch", config={"displayModeBar": False})
            with d2:
                st.markdown(leyenda_html(rep["dist"], rep["total"]), unsafe_allow_html=True)
        st.markdown('<div class="aviso">Las cifras de desempeño del modelo (precisión, recall, F1) se documentan en el informe de '
                    'validación del proyecto. Elaborado por <b>' + html.escape(autor) + '</b>.</div>', unsafe_allow_html=True)


COLOR_NIVEL = {"BAJO": "#3C8C40", "MEDIO": "#E0A93B", "ALTO": "#E23D3D"}
COLOR_GRAVEDAD = {"alta": "#E23D3D", "media": "#E0A93B", "baja": "#6B7B8C"}
QUE_HACER = {
    "ALTO": "No abras el enlace ni escribas contraseñas o datos personales. Si te llegó por correo o mensaje, no respondas y avísale a tu equipo de TI. "
            "Si ya lo abriste e ingresaste datos, cambia tus contraseñas de inmediato.",
    "MEDIO": "Ten cuidado. No ingreses datos desde este enlace: escribe tú mismo en el navegador la dirección oficial de la entidad y compárala.",
    "BAJO": "Esto no es una garantía: si el mensaje te pide datos o dinero con urgencia, desconfía igual y, ante la duda, entra escribiendo tú la dirección oficial.",
    "ACORTADO": "Los enlaces acortados esconden el destino y este análisis no puede evaluarlo. Pide el enlace completo a quien lo envió antes de abrirlo.",
}


def poner_ejemplo(texto):
    st.session_state["enlaces_txt"] = texto


def tarjeta_enlace(r):
    if r.get("error"):
        return ('<div class="enlace-card" style="border-left-color:#6B7B8C;"><div class="enlace-url">' + html.escape(r["url"][:120])
                + '</div><div class="razon"><span class="pt" style="background:#E23D3D"></span><span>' + html.escape(r["error"]) + '</span></div></div>')
    nivel = r["nivel"]
    color = COLOR_NIVEL[nivel]
    url = r["url"] if len(r["url"]) <= 110 else r["url"][:110] + "…"
    if r["acortado"]:
        cabecera = chip("RIESGO " + nivel + " · NO EVALUABLE", SEV_ESTILO[nivel], True)
        barra = ""
        consejo = QUE_HACER["ACORTADO"]
    else:
        pct = round(r["p"] * 100)
        cabecera = chip("RIESGO " + nivel, SEV_ESTILO[nivel], True) + ' &nbsp;<b style="font-size:13px;">Probabilidad de phishing: ' + str(pct) + '%</b>'
        barra = '<div class="barra"><div style="width:' + str(max(pct, 2)) + '%;background:' + color + ';"></div></div>'
        consejo = QUE_HACER[nivel]
    razones = r["razones"]
    if razones:
        lista = "".join('<div class="razon"><span class="pt" style="background:' + COLOR_GRAVEDAD[g] + '"></span><span>' + html.escape(t) + '</span></div>' for g, t in razones)
    elif nivel == "BAJO":
        lista = '<div class="razon"><span class="pt" style="background:#3C8C40"></span><span>No se encontraron señales de riesgo en el texto del enlace.</span></div>'
    else:
        lista = ('<div class="razon"><span class="pt" style="background:#E0A93B"></span><span>La IA encontró una combinación de rasgos parecida a la de '
                 'enlaces de phishing conocidos, sin una señal puntual que se pueda explicar con una sola regla.</span></div>')
    return ('<div class="enlace-card" style="border-left-color:' + color + ';">' + cabecera + barra
            + '<div class="enlace-url">' + html.escape(url) + '</div><div class="panel-sub" style="margin-bottom:2px;"><b>Por qué</b></div>'
            + lista + '<div class="que-hacer"><b>Qué hacer:</b> ' + html.escape(consejo) + '</div></div>')


def vista_enlaces(usuario):
    st.markdown('<div class="dash-title">◈ Verificador de enlaces</div>'
                '<div class="dash-sub">Pega un enlace sospechoso (por ejemplo, de un correo o un mensaje) y Centinela te dice qué tan probable es que sea phishing y por qué. '
                'No abre el enlace ni lo envía a ningún servicio externo.</div><br>', unsafe_allow_html=True)
    if ph is None or not (os.path.exists("modelo_phishing.json") and os.path.exists("phishing_meta.json")):
        st.error("Faltan archivos del verificador (phishing.py, modelo_phishing.json y phishing_meta.json) en la carpeta del proyecto. "
                 "Cópialos junto a app.py y recarga.")
        return
    try:
        modelo, meta = cargar_phishing()
    except Exception as e:
        st.error("No pude cargar el modelo de phishing: " + str(e)[:200])
        return

    with st.container(key="card_enlaces_input"):
        st.markdown('<div class="panel-title">Enlaces a revisar</div>'
                    '<div class="panel-sub">Uno por línea, hasta ' + str(MAX_ENLACES) + '. Puedes copiar los enlaces de un correo sospechoso.</div>',
                    unsafe_allow_html=True)
        st.text_area("Enlaces", key="enlaces_txt", height=110, label_visibility="collapsed",
                     placeholder="https://ejemplo.com/ruta\nsitio-sospechoso.xyz/login")
        analizar = st.button("Analizar enlaces", key="btn_enlaces", type="primary")

    if analizar:
        lineas = [l for l in st.session_state.get("enlaces_txt", "").splitlines() if l.strip()]
        if not lineas:
            st.warning("Pega al menos un enlace.")
        else:
            if len(lineas) > MAX_ENLACES:
                st.warning("Máximo " + str(MAX_ENLACES) + " enlaces por análisis: se revisaron los primeros " + str(MAX_ENLACES) + ".")
                lineas = lineas[:MAX_ENLACES]
            res = analizar_enlaces(lineas, modelo)
            st.session_state["enlaces_res"] = res
            hist = st.session_state.setdefault("enlaces_hist", [])
            hora = datetime.now().strftime("%H:%M:%S")
            for r in res:
                if not r.get("error"):
                    hist.insert(0, {"hora": hora, "url": r["url"], "nivel": r["nivel"], "p": r["p"]})
            st.session_state["enlaces_hist"] = hist[:10]

    res = st.session_state.get("enlaces_res")
    if res:
        validos = [r for r in res if not r.get("error")]
        conteo = {n: sum(1 for r in validos if r["nivel"] == n) for n in ("ALTO", "MEDIO", "BAJO")}
        st.markdown('<div class="panel-title" style="margin:6px 0 10px 0;">Resultado: ' + str(len(validos)) + ' enlace(s) revisado(s) · '
                    + str(conteo["ALTO"]) + ' de riesgo alto · ' + str(conteo["MEDIO"]) + ' medio · ' + str(conteo["BAJO"]) + ' bajo</div>',
                    unsafe_allow_html=True)
        for r in sorted(res, key=lambda x: ({"ALTO": 0, "MEDIO": 1, "BAJO": 2}.get(x.get("nivel"), 3))):
            st.markdown(tarjeta_enlace(r), unsafe_allow_html=True)

    hist = st.session_state.get("enlaces_hist", [])
    if hist:
        with st.container(key="card_enlaces_hist"):
            filas = "".join('<tr><td class="mono">' + h["hora"] + '</td><td class="mono">' + html.escape(h["url"][:70]) + '</td><td>'
                            + chip(h["nivel"], SEV_ESTILO[h["nivel"]], True) + '</td><td>' + (str(round(h["p"] * 100)) + '%' if h["p"] is not None else 'No evaluable')
                            + '</td></tr>' for h in hist)
            st.markdown('<div class="panel-title">Últimas verificaciones de esta sesión</div><div class="tabla-wrap"><table class="tabla-alertas"><thead><tr>'
                        '<th>HORA</th><th>ENLACE</th><th>RIESGO</th><th>PROBABILIDAD</th></tr></thead><tbody>' + filas + '</tbody></table></div>',
                        unsafe_allow_html=True)

    m = meta.get("metricas", {})
    exact = m.get("v3_global", {}).get("exactitud")
    mod = m.get("por_grupo_marcados_phishing", {}).get("moderno_malo", {}).get("v3")
    col = m.get("por_grupo_marcados_phishing", {}).get("colombia_bueno", {}).get("v3")
    with st.expander("¿Cómo funciona y cuáles son sus límites?"):
        cifras = ""
        if exact is not None and mod is not None and col is not None:
            cifras = ('<li>En pruebas con sitios que el modelo nunca vio: exactitud de <b>' + "{:.1f}".format(exact * 100) + '%</b>, detecta <b>'
                      + "{:.1f}".format(mod * 100) + '%</b> del phishing reciente y marca por error <b>' + "{:.1f}".format(col * 100)
                      + '%</b> de sitios colombianos legítimos.</li>')
        st.markdown('<ul style="font-size:13px;line-height:1.6;">'
                    '<li>Solo analiza el <b>texto del enlace</b> (largo, palabras, dominio, símbolos). No abre el sitio, no mira su contenido ni su certificado.</li>'
                    '<li>El análisis ocurre en este equipo: el enlace no se envía a ningún servicio externo.</li>' + cifras +
                    '<li>Páginas legítimas de inicio de sesión pueden salir con riesgo medio o alto, porque el texto del enlace se parece al de un engaño.</li>'
                    '<li>Un enlace acortado (bit.ly y similares) no se puede evaluar porque esconde su destino.</li>'
                    '<li>No reemplaza el criterio ni los filtros de correo: es una capa de apoyo. Ante la duda, entra escribiendo tú la dirección oficial.</li>'
                    '</ul>', unsafe_allow_html=True)


QUE_HACER_CORREO = {
    "ALTO": "No hagas clic en los enlaces ni abras los adjuntos, y no respondas. Avísale a tu equipo de TI y, si ya entregaste datos, cambia tus contraseñas de inmediato.",
    "MEDIO": "Desconfía. Antes de actuar, verifica por otro canal (llama a un número oficial o escribe tú la dirección en el navegador). Si no estás segura, trátalo como phishing.",
    "BAJO": "No se encontraron señales claras de phishing. Eso no es una garantía: si el mensaje pide datos, dinero o urgencia, desconfía igual.",
}
CHIP_AUTH = {"pass": ("#2E7D32", "#E3F4E5"), "fail": ("#C62828", "#FDE8E8"), "softfail": ("#946200", "#FFF4D6"), "permerror": ("#C62828", "#FDE8E8")}


def poner_ejemplo_correo(nombre):
    st.session_state["correo_txt"] = cx.ejemplos()[nombre].decode("utf-8", "ignore")


def _pct(x):
    return "{:.0f}%".format(x * 100) if x is not None else "no evaluado"


def tarjeta_correo(r):
    if r.get("error"):
        return ('<div class="enlace-card" style="border-left-color:#6B7B8C;"><b>' + html.escape(r["nombre"][:80]) + '</b>'
                '<div class="razon"><span class="pt" style="background:#E23D3D"></span><span>' + html.escape(r["error"]) + '</span></div></div>')
    nivel, pct = r["nivel"], round(r["p_final"] * 100)
    color = COLOR_NIVEL[nivel]
    de = html.escape(((r["de_nombre"] + " ") if r["de_nombre"] else "") + ("<" + r["de_correo"] + ">" if r["de_correo"] else "")) or "remitente desconocido"
    idioma = {"es": "español", "en": "inglés", "pt": "portugués"}.get(r["idioma"], "no identificado")
    razones = "".join('<div class="razon"><span class="pt" style="background:' + COLOR_GRAVEDAD[x["gravedad"]] + '"></span><span class="area">' + html.escape(x["area"])
                      + '</span><span>' + html.escape(x["texto"]) + '</span></div>' for x in r["razones"])
    if not razones:
        razones = '<div class="razon"><span class="pt" style="background:#3C8C40"></span><span>No se encontraron señales de riesgo.</span></div>'
    positivas = "".join('<div class="razon"><span class="pt" style="background:#3C8C40"></span><span class="area">A favor</span><span>' + html.escape(t) + '</span></div>' for t in r["positivas"])
    aviso = ""
    if r["modelo_texto"] is None:
        aviso = ('<div class="aviso-amarillo">No hay modelo de texto para el idioma de este correo (' + idioma + '): el resultado se basa solo en señales técnicas y enlaces, '
                 'y puede ser menos confiable.</div>')
    inc = ('<div class="aviso" style="margin-top:8px;">Se registró como <b>incidente #' + str(r["incidente"]) + '</b> en el Historial.</div>') if r.get("incidente") else ""
    return ('<div class="enlace-card" style="border-left-color:' + color + ';">' + chip("RIESGO " + nivel, SEV_ESTILO[nivel], True)
            + ' &nbsp;<b style="font-size:13px;">Probabilidad de phishing: ' + str(pct) + '%</b>'
            + '<div class="barra"><div style="width:' + str(max(pct, 2)) + '%;background:' + color + ';"></div></div>'
            + '<div class="panel-title" style="margin-top:8px;">' + html.escape(r["asunto"][:110] or "(sin asunto)") + '</div>'
            + '<div class="panel-sub">De: ' + de + ' · Idioma: ' + idioma + '</div>'
            + '<div class="panel-sub" style="margin:8px 0 2px 0;"><b>Por qué</b></div>' + razones + positivas + aviso
            + '<div class="que-hacer"><b>Qué hacer:</b> ' + html.escape(QUE_HACER_CORREO[nivel]) + '</div>' + inc + '</div>')


def detalle_correo_html(r):
    def chip_auth(v):
        fg, bg = CHIP_AUTH.get(v, ("#475569", "#E8EDF3"))
        return '<span class="chip" style="color:' + fg + ';background:' + bg + ';">' + html.escape(v or "sin dato") + '</span>'
    comp = ('<tr><td>Texto del mensaje</td><td>' + (_pct(r["p_texto"]) if r["modelo_texto"] else "no evaluado") + '</td></tr>'
            '<tr><td>Enlaces (apoyo)</td><td>' + (_pct(r["p_enlaces"]) if r["p_enlaces"] is not None else "sin enlaces evaluables") + '</td></tr>'
            '<tr><td>Señales técnicas</td><td>' + (_pct(r["p_tecnico"]) if r["p_tecnico"] is not None else "ninguna") + '</td></tr>'
            '<tr><td><b>Resultado combinado</b></td><td><b>' + _pct(r["p_final"]) + '</b></td></tr>')
    auth = ('SPF ' + chip_auth(r["auth"]["spf"]) + ' &nbsp; DKIM ' + chip_auth(r["auth"]["dkim"]) + ' &nbsp; DMARC ' + chip_auth(r["auth"]["dmarc"])) if r["tiene_auth"] else \
        'No hay encabezados de autenticación en este correo (por ejemplo, si se copió solo el texto): no se pudo verificar SPF, DKIM ni DMARC.'
    enl = "".join('<tr><td class="mono">' + html.escape(e["url"][:70]) + '</td><td>' + html.escape((e["ancla"] or "")[:30]) + '</td><td>'
                  + ("acortado: no evaluable" if e["acortado"] else _pct(e["p"])) + '</td></tr>' for e in r["enlaces"]) or '<tr><td colspan="3">Sin enlaces.</td></tr>'
    adj = "".join('<li>' + html.escape(a["nombre"]) + ' <small>(' + html.escape(a["tipo"]) + ', ' + str(a["bytes"]) + ' bytes)</small></li>' for a in r["adjuntos"]) or "<li>Sin adjuntos.</li>"
    pal = ", ".join('«' + html.escape(w) + '»' for w, _ in r["palabras"]) or "ninguna destacada"
    return ('<div class="panel-sub"><b>Cómo se combinó el resultado</b></div><table class="tabla-alertas"><tbody>' + comp + '</tbody></table>'
            '<div class="panel-sub" style="margin-top:10px;"><b>Autenticación del remitente</b></div><div style="font-size:13px;">' + auth + '</div>'
            '<div class="panel-sub" style="margin-top:10px;"><b>Enlaces del correo</b> (el riesgo por enlace es orientativo)</div>'
            '<table class="tabla-alertas"><thead><tr><th>ENLACE</th><th>TEXTO QUE SE VE</th><th>RIESGO</th></tr></thead><tbody>' + enl + '</tbody></table>'
            '<div class="panel-sub" style="margin-top:10px;"><b>Adjuntos</b></div><ul style="font-size:13px;margin:0;">' + adj + '</ul>'
            '<div class="panel-sub" style="margin-top:10px;"><b>Palabras del mensaje que más pesaron</b> (aporte exacto del modelo): ' + pal + '</div>')


def vista_correos(usuario):
    st.markdown('<div class="dash-title">◆ Analizar correos</div>'
                '<div class="dash-sub">Sube un correo sospechoso (archivo .eml) o pega su contenido completo. Centinela revisa el remitente, la autenticación (SPF, DKIM y DMARC), '
                'los enlaces, los adjuntos y el texto, y te dice qué tan probable es que sea phishing y por qué. Todo se analiza en este equipo y solo se guardan '
                'los datos que extrae (remitente, asunto y razones), nunca el mensaje.</div><br>', unsafe_allow_html=True)
    faltan = [f for f in ("modelo_texto_es.json", "modelo_texto_en.json", "modelo_enlaces_correo.json", "phishing.py", "correos.py") if not os.path.exists(f)]
    if cx is None or ph is None or faltan:
        st.error("Faltan archivos del analizador de correos en la carpeta del proyecto: " + (", ".join(faltan) or "correos.py o phishing.py") + ". Cópialos junto a app.py y recarga.")
        return
    try:
        modelos = cargar_correos()
    except Exception as e:
        st.error("No pude cargar los modelos del analizador: " + str(e)[:200])
        return
    with st.container(key="card_correos_input"):
        st.markdown('<div class="panel-title">Correos a revisar</div>'
                    '<div class="panel-sub">Archivos .eml (hasta ' + str(MAX_CORREOS) + ', 5 MB cada uno). En Gmail: abre el correo, tres puntos y «Descargar mensaje». '
                    'No abras los correos sospechosos con tu programa de correo.</div>', unsafe_allow_html=True)
        archivos = st.file_uploader("Correos", type=["eml", "txt"], accept_multiple_files=True, key="correo_files", label_visibility="collapsed")
        st.text_area("Correo pegado", key="correo_txt", height=120, label_visibility="collapsed",
                     placeholder="O pega aquí el correo completo (con los encabezados: From, Subject, Authentication-Results...)")
        analizar = st.button("Analizar correos", key="btn_correos", type="primary")

    if analizar:
        entradas = []
        for f in (archivos or [])[:MAX_CORREOS]:
            datos = f.getvalue()
            entradas.append((f.name, None if len(datos) > MAX_BYTES_CORREO else datos, "El archivo pesa más de 5 MB." if len(datos) > MAX_BYTES_CORREO else None))
        if archivos and len(archivos) > MAX_CORREOS:
            st.warning("Máximo " + str(MAX_CORREOS) + " archivos por análisis: se revisaron los primeros " + str(MAX_CORREOS) + ".")
        pegado = st.session_state.get("correo_txt", "").strip()
        if pegado:
            entradas.append(("Correo pegado", None if len(pegado) > 2_000_000 else pegado.encode("utf-8", "ignore"), "El texto pegado es demasiado largo." if len(pegado) > 2_000_000 else None))
        if not entradas:
            st.warning("Sube un archivo .eml o pega el contenido de un correo.")
        else:
            resultados = []
            for nombre, datos, error in entradas:
                if error:
                    resultados.append({"nombre": nombre, "error": error})
                    continue
                try:
                    r = cx.analizar_correo(datos, modelos)
                    if r["largo_cuerpo"] < 3 and not r["de_correo"] and not r["asunto"]:
                        raise ValueError("no encontré remitente, asunto ni texto: no parece un correo.")
                    r["nombre"] = nombre
                    resultados.append(r)
                except Exception as e:
                    resultados.append({"nombre": nombre, "error": "No pude leer este correo: " + str(e)[:120]})
            validos = [r for r in resultados if not r.get("error")]
            ids = guardar_analisis_correos(usuario["id"], validos)[1] if validos else []
            it = iter(ids)
            for r in validos:
                r["incidente"] = next(it) if r["nivel"] in ("ALTO", "MEDIO") else None
            st.session_state["correos_res"] = resultados

    res = st.session_state.get("correos_res")
    if res:
        validos = [r for r in res if not r.get("error")]
        cnt = {n: sum(1 for r in validos if r["nivel"] == n) for n in ("ALTO", "MEDIO", "BAJO")}
        st.markdown('<div class="panel-title" style="margin:6px 0 10px 0;">Resultado: ' + str(len(validos)) + ' correo(s) revisado(s) · ' + str(cnt["ALTO"]) + ' de riesgo alto · '
                    + str(cnt["MEDIO"]) + ' medio · ' + str(cnt["BAJO"]) + ' bajo</div>', unsafe_allow_html=True)
        for r in sorted(res, key=lambda x: -x.get("p_final", -1)):
            st.markdown(tarjeta_correo(r), unsafe_allow_html=True)
            if not r.get("error"):
                with st.expander("Ver detalle técnico de «" + (r["asunto"][:50] or r["nombre"][:50]) + "»"):
                    st.markdown(detalle_correo_html(r), unsafe_allow_html=True)
        if any(r.get("incidente") for r in validos):
            st.button("Ver los incidentes en el Historial ›", key="btn_correos_hist", type="tertiary", on_click=ir_seccion, args=("historial",))

    with st.expander("¿Cómo funciona y cuáles son sus límites?"):
        st.markdown('<ul style="font-size:13px;line-height:1.6;">'
                    '<li><b>Qué combina:</b> el texto del mensaje (modelo en español o en inglés), los enlaces (como apoyo), y señales técnicas explicables: autenticación SPF, DKIM y DMARC, remitente, '
                    'enlaces engañosos y adjuntos peligrosos.</li>'
                    '<li><b>Qué tan confiable es el texto:</b> el modelo en español, entrenado con SpaPhish, llega a un F1 de 0,95 con correos de la misma fuente, pero baja a cerca de 0,86 al probarlo con '
                    'correos más recientes (detecta unos 3 de cada 4 phishing nuevos). El modelo en inglés, entrenado con corpus antiguos, no sirve para español. '
                    'Aún no se ha medido con correos colombianos.</li>'
                    '<li><b>Los enlaces apoyan, no mandan:</b> al evaluarlos con separación estricta por dominio, su capacidad de detección fue baja, así que por sí solos solo pueden subir el riesgo a medio.</li>'
                    '<li><b>Las señales técnicas</b> usan pesos definidos por criterio de experto, no aprendidos con datos; no se pudieron medir con correos legítimos que traigan encabezados.</li>'
                    '<li>Con idiomas sin modelo (por ejemplo, portugués) el análisis se basa solo en señales técnicas y es menos confiable.</li>'
                    '<li>Un correo de un remitente real que fue hackeado puede pasar SPF, DKIM y DMARC. No reemplaza el criterio ni los filtros de tu proveedor de correo: es una capa de apoyo.</li>'
                    '</ul>', unsafe_allow_html=True)


def _corto(t, n):
    t = str(t)
    return t if len(t) <= n else t[: n - 1] + "…"


GUIA_CONECTOR_HTML = (
    '<ol style="font-size:13px;line-height:1.75;padding-left:20px;margin:0;">'
    '<li><b>Ten a la mano la cuenta de Gmail de tu empresa.</b> Debe ser una cuenta propia o autorizada por su dueño.</li>'
    '<li><b>Activa la verificación en dos pasos.</b> Entra a <b>myaccount.google.com</b>, abre <b>Seguridad</b> y activa «Verificación en dos pasos» (te pide un celular). Sin esto Google no deja crear el siguiente paso.</li>'
    '<li><b>Crea una contraseña de aplicación.</b> En myaccount.google.com escribe «Contraseñas de aplicaciones» en el buscador (o entra a <b>myaccount.google.com/apppasswords</b>). Ponle de nombre «Centinela», dale <b>Crear</b> y copia el código de 16 letras: Google lo muestra una sola vez.</li>'
    '<li><b>Revisa que IMAP esté activado</b> (solo si aparece la opción). En Gmail: engranaje, «Ver toda la configuración», pestaña «Reenvío y correo POP/IMAP», «Habilitar IMAP», guardar.</li>'
    '<li><b>Conecta aquí.</b> Escribe el correo y pega el código de 16 letras (con o sin espacios). Dale <b>«1. Probar conexión»</b>: debe salir un cuadro verde que dice «Conexión correcta».</li>'
    '<li><b>Elige qué hacer con los correos de riesgo alto.</b> «Solo avisar» crea el incidente y no toca tu bandeja (recomendado para empezar). «Cuarentena automática» además los manda a la etiqueta «Centinela/Cuarentena».</li>'
    '<li><b>Revisa el buzón.</b> «2. Revisar el buzón ahora» lo hace una vez; «Vigilar cada 30 segundos» lo repite solo mientras esta pantalla esté abierta. Para devolver un correo de la cuarentena, abre el incidente en el Historial y usa «Liberar».</li>'
    '<li><b>Recibe avisos en tu correo.</b> Con la opción «Avisarme por correo» marcada, Centinela te envía a tu Gmail una advertencia cuando detecta algo sospechoso: con el logo, el tipo de amenaza, por qué y qué hacer. Llega sola, mientras la vigilancia esté activa.</li>'
    '</ol>'
    '<div class="aviso" style="margin-top:10px;"><b>Si algo falla:</b> «Gmail rechazó el correo o la contraseña» significa que el código está mal escrito o que usaste la contraseña normal: crea otra contraseña de aplicación. '
    'Si no aparece «Contraseñas de aplicaciones», la verificación en dos pasos no está activa (o la cuenta es de trabajo y el administrador la bloquea). '
    'Para quitarle el acceso a Centinela, borra la contraseña en myaccount.google.com/apppasswords (bote de basura).</div>'
    '<div class="aviso-amarillo"><b>Privacidad:</b> la contraseña solo vive en la memoria de esta sesión. Centinela guarda del correo el remitente, el asunto, las razones y los puntajes, nunca el contenido. '
    'Conecta únicamente buzones propios o con permiso de su dueño.</div>')
MAX_POR_REVISION = 20
INTERVALO_BUZON = 30


def alternar_vigilancia():
    if st.session_state.get("buzon_vigilando"):
        st.session_state["buzon_vigilando"] = False
        return
    u, c = st.session_state.get("buzon_usuario", "").strip(), st.session_state.get("buzon_clave_txt", "")
    if not u or not c.strip():
        st.session_state["toast"] = "Escribe el correo y la contraseña de aplicación antes de vigilar."
        return
    st.session_state["buzon_cfg"] = {"usuario": u, "clave": c}
    st.session_state["buzon_vigilando"] = True
    st.session_state["buzon_ultimo_chequeo"] = 0.0


def fijar_cfg_desde_campos():
    u, c = st.session_state.get("buzon_usuario", "").strip(), st.session_state.get("buzon_clave_txt", "")
    if u and c.strip():
        st.session_state["buzon_cfg"] = {"usuario": u, "clave": c}
    return st.session_state.get("buzon_cfg")


def _destino_aviso(cfg):
    d = st.session_state.get("buzon_destino", "").strip()
    return d if re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", d) else cfg["usuario"]


def _logo_aviso():
    for f in ("logo_centinela.png", "isotipo_centinela.png"):
        if os.path.exists(f):
            with open(f, "rb") as fh:
                return fh.read()
    return None


def enviar_aviso_buzon(cfg, encontrados):
    """Escribe un solo correo de aviso con todos los sospechosos de esta revision (si la opcion esta activa)."""
    if not st.session_state.get("buzon_avisar", True):
        return
    niveles = ("ALTO",) if st.session_state.get("buzon_nivel_aviso", "Solo riesgo alto") == "Solo riesgo alto" else ("ALTO", "MEDIO")
    susp = []
    for x in encontrados:
        r = x["r"]
        if r is None or r["nivel"] not in niveles:
            continue
        graves = [q["texto"] for q in r["razones"] if q["gravedad"] in ("alta", "media")] or [q["texto"] for q in r["razones"]]
        susp.append({"nivel": r["nivel"], "de": r["de_correo"] or r["de_nombre"] or "remitente desconocido", "asunto": r["asunto"] or "(sin asunto)",
                     "tipos": imap_conector.tipos_de_amenaza(r), "razones": graves[:3],
                     "accion": ("Quedó en cuarentena: búscalo en la carpeta «Centinela/Cuarentena» de Gmail." if r.get("en_cuarentena")
                                else "Sigue en tu bandeja: no abras sus enlaces ni adjuntos.")})
    if not susp:
        return
    destino = _destino_aviso(cfg)
    try:
        imap_conector.enviar_aviso(cfg["usuario"], cfg["clave"], imap_conector.construir_aviso(cfg["usuario"], destino, susp, _logo_aviso()))
        st.session_state["buzon_aviso_estado"] = {"ok": True, "texto": "Se envió un aviso a " + destino + " con " + str(len(susp)) + " correo(s) sospechoso(s)."}
    except imap_conector.ErrorBuzon as e:
        st.session_state["buzon_aviso_estado"] = {"ok": False, "texto": str(e)}


def revisar_buzon(usuario, modelos):
    cfg = st.session_state.get("buzon_cfg")
    if not cfg:
        st.session_state["buzon_error"] = "Primero escribe el correo y la contraseña de aplicación."
        return 0
    ultimo = st.session_state.get("buzon_ultimo_uid")
    modo = st.session_state.get("buzon_modo", "Solo avisar")
    encontrados = []
    try:
        with imap_conector.BuzonIMAP(cfg["usuario"], cfg["clave"]) as buzon:
            for uid, raw in buzon.nuevos(ultimo, MAX_POR_REVISION if ultimo is not None else 10):
                ultimo = max(ultimo or 0, uid)
                if raw is None:
                    encontrados.append({"r": None, "accion": "Omitido: el correo es demasiado grande"})
                    continue
                if b"X-Centinela-Alerta:" in raw[:30000]:
                    marcar_procesado(usuario["id"], imap_conector.id_mensaje(raw))
                    continue
                try:
                    r = cx.analizar_correo(raw, modelos)
                except Exception:
                    encontrados.append({"r": None, "accion": "No pude leer este correo"})
                    continue
                mid = imap_conector.id_mensaje(raw)
                r["mensaje_id"] = mid
                if not marcar_procesado(usuario["id"], mid):
                    continue
                accion = "Sin acción"
                if r["nivel"] == "ALTO" and modo == "Cuarentena automática":
                    if mid.startswith("sha1:"):
                        accion = "Riesgo alto, pero sin identificador: solo avisó"
                    else:
                        try:
                            buzon.a_cuarentena(uid)
                            r["en_cuarentena"] = True
                            accion = "Movido a cuarentena"
                        except imap_conector.ErrorBuzon as e:
                            accion = "No se pudo mover: " + str(e)[:60]
                elif r["nivel"] in ("ALTO", "MEDIO"):
                    accion = "Solo avisó"
                encontrados.append({"r": r, "accion": accion})
    except imap_conector.ErrorBuzon as e:
        st.session_state["buzon_error"] = str(e)
        return 0
    st.session_state["buzon_error"] = None
    st.session_state["buzon_ultimo_uid"] = ultimo
    validos = [x["r"] for x in encontrados if x["r"] is not None]
    if validos:
        guardar_analisis_correos(usuario["id"], validos, origen="buzon")
    enviar_aviso_buzon(cfg, encontrados)
    stats = st.session_state.setdefault("buzon_stats", {"revisados": 0, "alertas": 0, "cuarentena": 0})
    log = st.session_state.setdefault("buzon_log", [])
    hora = datetime.now().strftime("%H:%M:%S")
    for x in encontrados:
        r = x["r"]
        if r is None:
            log.insert(0, {"hora": hora, "de": "-", "asunto": "-", "nivel": None, "accion": x["accion"]})
            continue
        stats["revisados"] += 1
        stats["alertas"] += 1 if r["nivel"] in ("ALTO", "MEDIO") else 0
        stats["cuarentena"] += 1 if r.get("en_cuarentena") else 0
        log.insert(0, {"hora": hora, "de": r["de_correo"] or r["de_nombre"] or "desconocido", "asunto": r["asunto"] or "(sin asunto)", "nivel": r["nivel"], "accion": x["accion"]})
    st.session_state["buzon_log"] = log[:15]
    st.session_state["buzon_ultimo_chequeo"] = time.time()
    return sum(1 for r in validos if r["nivel"] in ("ALTO", "MEDIO"))


def panel_buzon(usuario, modelos):
    if st.session_state.get("buzon_vigilando"):
        ahora_t = time.time()
        if ahora_t - st.session_state.get("buzon_ultimo_chequeo", 0.0) >= INTERVALO_BUZON - 5:
            n = revisar_buzon(usuario, modelos)
            if n and ahora_t - st.session_state.get("buzon_ultimo_rerun", 0.0) > 10:
                st.session_state["buzon_ultimo_rerun"] = ahora_t
                st.toast("🚨 " + str(n) + " correo(s) sospechoso(s) detectado(s)")
                st.rerun()
    if st.session_state.get("buzon_error"):
        st.error(st.session_state["buzon_error"])
    est_aviso = st.session_state.get("buzon_aviso_estado")
    if est_aviso:
        (st.success if est_aviso["ok"] else st.warning)(("✉ " if est_aviso["ok"] else "No se pudo enviar el aviso. ") + est_aviso["texto"])
    stats = st.session_state.get("buzon_stats", {"revisados": 0, "alertas": 0, "cuarentena": 0})
    ult = st.session_state.get("buzon_ultimo_chequeo", 0.0)
    ult_txt = datetime.fromtimestamp(ult).strftime("%H:%M:%S") if ult else "aún no"
    vig = st.session_state.get("buzon_vigilando")
    estado_html = ('<span class="chip" style="color:#2E7D32;background:#E3F4E5;"><span class="dot-live"></span>VIGILANDO</span>' if vig
                   else '<span class="chip" style="color:#475569;background:#E8EDF3;">EN PAUSA</span>')
    k1, k2, k3, k4 = st.columns(4)
    k1.markdown('<div class="kpi-card" style="min-height:118px"><div class="kpi-label-top">Vigilancia</div><div style="margin-top:10px;">' + estado_html
                + '</div><div class="kpi-sub">Última revisión: ' + ult_txt + '</div></div>', unsafe_allow_html=True)
    k2.markdown('<div class="kpi-card" style="min-height:118px"><div class="kpi-label-top">Correos revisados</div><div class="kpi-num-big">' + str(stats["revisados"]) + '</div></div>', unsafe_allow_html=True)
    k3.markdown('<div class="kpi-card" style="min-height:118px"><div class="kpi-label-top">Sospechosos</div><div class="kpi-num-big">' + str(stats["alertas"]) + '</div></div>', unsafe_allow_html=True)
    k4.markdown('<div class="kpi-card" style="min-height:118px"><div class="kpi-label-top">En cuarentena</div><div class="kpi-num-big" style="color:#6D28D9">' + str(stats["cuarentena"]) + '</div></div>', unsafe_allow_html=True)
    st.write("")
    log = st.session_state.get("buzon_log", [])
    with st.container(key="card_buzon_log"):
        st.markdown('<div class="panel-title">Correos revisados del buzón</div><div class="panel-sub">Los más recientes primero. Los sospechosos quedan como incidentes en el Historial.</div>', unsafe_allow_html=True)
        if not log:
            st.markdown('<div class="panel-sub">Todavía no se ha revisado ningún correo. Prueba la conexión y luego revisa el buzón.</div>', unsafe_allow_html=True)
        else:
            filas = "".join('<tr><td class="mono">' + x["hora"] + '</td><td>' + html.escape(_corto(x["de"], 38)) + '</td><td>' + html.escape(_corto(x["asunto"], 56)) + '</td><td>'
                            + (chip(x["nivel"], SEV_ESTILO[x["nivel"]], True) if x["nivel"] else "") + '</td><td>' + html.escape(x["accion"]) + '</td></tr>' for x in log)
            st.markdown('<div class="tabla-wrap"><table class="tabla-alertas"><thead><tr><th>HORA</th><th>DE</th><th>ASUNTO</th><th>RIESGO</th><th>QUÉ HIZO CENTINELA</th></tr></thead><tbody>'
                        + filas + '</tbody></table></div>', unsafe_allow_html=True)
            if st.button("Ver los incidentes en el Historial ›", key="btn_buzon_hist", type="tertiary"):
                ir_seccion("historial")
                st.rerun()


def vista_conector(usuario):
    st.markdown('<div class="dash-title">⇄ Conector de correo</div>'
                '<div class="dash-sub">Agrega la cuenta de Gmail de tu empresa y Centinela revisa los correos que llegan, calcula el riesgo de phishing y, si es alto, '
                'lo manda a la carpeta «Centinela/Cuarentena» del buzón. <b>Nunca borra correos.</b></div><br>', unsafe_allow_html=True)
    faltan = [f for f in ("conector.py", "correos.py", "phishing.py", "modelo_texto_es.json", "modelo_texto_en.json", "modelo_enlaces_correo.json") if not os.path.exists(f)]
    if imap_conector is None or cx is None or ph is None or faltan:
        st.error("Faltan archivos del conector en la carpeta del proyecto: " + (", ".join(faltan) or "conector.py") + ". Cópialos junto a app.py y recarga.")
        return
    try:
        modelos = cargar_correos()
    except Exception as e:
        st.error("No pude cargar los modelos del analizador: " + str(e)[:200])
        return
    with st.expander("¿Primera vez? Guía paso a paso para conectar tu Gmail", expanded="buzon_cfg" not in st.session_state):
        st.markdown(GUIA_CONECTOR_HTML, unsafe_allow_html=True)
    with st.container(key="card_conector_cfg"):
        st.markdown('<div class="panel-title">Conexión con el buzón</div>'
                    '<div class="panel-sub">Agrega tu cuenta de Gmail y su <b>contraseña de aplicación</b>: el código de 16 letras que genera Google '
                    '(no es la contraseña normal de tu correo). Vive solo en la memoria de esta sesión: no se guarda.</div>', unsafe_allow_html=True)
        c1, c2 = st.columns(2)
        c1.markdown('<div class="field-label">Cuenta de Gmail</div>', unsafe_allow_html=True)
        c1.text_input("Cuenta de Gmail", key="buzon_usuario", placeholder="tuempresa@gmail.com", label_visibility="collapsed")
        c2.markdown('<div class="field-label">Contraseña de aplicación (16 letras)</div>', unsafe_allow_html=True)
        c2.text_input("Contraseña de aplicación (16 letras)", key="buzon_clave_txt", type="password", placeholder="xxxx xxxx xxxx xxxx", label_visibility="collapsed")
        st.radio("¿Qué hace Centinela con los correos de riesgo alto?", ["Solo avisar", "Cuarentena automática"], key="buzon_modo", horizontal=True)
        st.caption("Con «Solo avisar» Centinela crea el incidente y no toca tu bandeja. La «Cuarentena automática» además mueve los correos de riesgo alto a la carpeta «Centinela/Cuarentena» (no los borra).")
        st.checkbox("Avisarme por correo cuando Centinela detecte algo sospechoso", value=True, key="buzon_avisar")
        a1, a2 = st.columns(2)
        a1.radio("Avisar cuando el riesgo sea", ["Solo riesgo alto", "Riesgo alto y medio"], key="buzon_nivel_aviso", horizontal=True)
        a2.markdown('<div class="field-label" style="margin-top:0;">Enviar los avisos a (opcional)</div>', unsafe_allow_html=True)
        a2.text_input("Enviar los avisos a (opcional)", key="buzon_destino", placeholder="Si lo dejas vacío, llegan a la misma cuenta", label_visibility="collapsed")
        b1, b2, b3 = st.columns(3)
        probar = b1.button("1. Probar conexión", key="btn_buzon_probar", type="primary", width="stretch")
        ahora = b2.button("2. Revisar el buzón ahora", key="btn_buzon_revisar", width="stretch")
        vig = st.session_state.get("buzon_vigilando")
        b3.button("■ Dejar de vigilar" if vig else "▶ Vigilar cada 30 segundos", key="btn_buzon_vigilar", on_click=alternar_vigilancia, width="stretch")
    if probar:
        cfg = fijar_cfg_desde_campos()
        if not cfg:
            st.warning("Escribe el correo y la contraseña de aplicación.")
        else:
            try:
                with imap_conector.BuzonIMAP(cfg["usuario"], cfg["clave"]) as buzon:
                    n = buzon.probar()
                st.success("Conexión correcta. La bandeja de entrada tiene " + str(n) + " correo(s).")
            except imap_conector.ErrorBuzon as e:
                st.error(str(e))
    if ahora:
        if fijar_cfg_desde_campos():
            revisar_buzon(usuario, modelos)
        else:
            st.warning("Escribe el correo y la contraseña de aplicación.")
    fragmento = st.fragment(run_every=str(INTERVALO_BUZON) + "s" if st.session_state.get("buzon_vigilando") else None)(panel_buzon)
    fragmento(usuario, modelos)
    with st.expander("¿Cómo funciona y cuáles son sus límites?"):
        st.markdown('<ul style="font-size:13px;line-height:1.6;">'
                    '<li>Centinela lee la bandeja de entrada sin marcar los correos como leídos, los analiza (texto, enlaces, autenticación, adjuntos) y los clasifica en riesgo bajo, medio o alto.</li>'
                    '<li>En modo <b>cuarentena automática</b>, los de riesgo alto se copian a «Centinela/Cuarentena» y se sacan de la bandeja. En Gmail esto solo quita la etiqueta de la bandeja de entrada: el correo sigue guardado. '
                    'Para devolverlo, abre el incidente en el Historial y usa «Liberar».</li>'
                    '<li>Revisa el buzón mientras esta pantalla esté abierta (cada 30 segundos). Un servicio que corra siempre en segundo plano queda como trabajo futuro.</li>'
                    '<li>La contraseña de aplicación solo vive en la memoria de la sesión. Si la ves en pantalla o en un archivo, bórrala desde tu cuenta de Google y crea otra.</li>'
                    '<li>Los riesgos de este analizador (idioma, remitentes hackeados, falsas alarmas) son los mismos de «Analizar correos». Una falsa alarma deja el correo en cuarentena hasta que lo liberes.</li>'
                    '</ul>', unsafe_allow_html=True)


def iniciar_sensor(usuario_id):
    st.session_state["sensor"] = nuevo_estado_sensor()
    st.session_state["sensor"]["activo"] = True
    st.session_state["sensor"]["sesion_id"] = iniciar_sesion_sensor(usuario_id)


def detener_sensor():
    if "sensor" in st.session_state:
        st.session_state["sensor"]["activo"] = False


def panel_sensor(usuario, grupos):
    s = st.session_state["sensor"]
    if s["activo"]:
        ahora_t = time.time()
        if ahora_t - s["ultimo_tick"] >= 1.2:
            s["ultimo_tick"] = ahora_t
            df = generar_lote(s, grupos, st.session_state.get("sensor_intensidad", "Media"),
                              VELOCIDAD_SENSOR[st.session_state.get("sensor_ritmo", "Media")],
                              st.session_state.get("sensor_tipos", CATEGORIAS))
            cats, certezas = clasificar_df(df)
            incs, conteo_lote = aplicar_lote(s, df, cats, certezas)
            registrar_lote_sensor(s["sesion_id"], usuario["id"], conteo_lote, incs)
            graves = [i for i in incs if i["severidad"] in ("CRÍTICO", "ALTO")]
            if graves and ahora_t - s["ultimo_toast"] > 4:
                s["ultimo_toast"] = ahora_t
                st.toast("🚨 " + str(len(incs)) + " alerta(s) nueva(s): " + NOMBRES_AMENAZA[graves[0]["categoria"]])
            if incs and ahora_t - s["ultimo_rerun"] > 6:
                s["ultimo_rerun"] = ahora_t
                st.rerun()

    if s["sesion_id"] is None:
        st.markdown('<div class="aviso">El sensor está apagado. Presiona <b>Iniciar monitoreo</b> y Centinela empezará a analizar '
                    'el tráfico simulado y a crear los incidentes por sí solo.</div>', unsafe_allow_html=True)
        return

    estado_html = ('<span class="chip" style="color:#2E7D32;background:#E3F4E5;"><span class="dot-live"></span>ACTIVO</span>'
                   if s["activo"] else '<span class="chip" style="color:#475569;background:#E8EDF3;">DETENIDO</span>')
    ultimo = s["feed"][-1]["hora"] if s["feed"] else "—"
    fg_c, bg_c = SEV_ESTILO["CRÍTICO"]
    k1, k2, k3, k4 = st.columns(4)
    k1.markdown('<div class="kpi-card" style="min-height:118px"><div class="kpi-label-top">Sensor virtual</div>'
                '<div style="margin-top:10px;">' + estado_html + '</div><div class="kpi-sub">Último evento: ' + ultimo + '</div></div>',
                unsafe_allow_html=True)
    k2.markdown('<div class="kpi-card" style="min-height:118px"><div class="kpi-label-top">Eventos analizados</div>'
                '<div class="kpi-num-big">' + str(s["total"]) + '</div><div class="kpi-sub">en esta sesión</div></div>',
                unsafe_allow_html=True)
    k3.markdown('<div class="kpi-card" style="min-height:118px"><div class="kpi-label-top">Alertas creadas</div>'
                '<div class="kpi-num-big">' + str(s["alertas_sesion"]) + '</div><div class="kpi-sub">incidentes nuevos en el historial</div></div>',
                unsafe_allow_html=True)
    k4.markdown('<div class="kpi-card" style="min-height:118px"><div class="kpi-top"><span class="kpi-label-top">Críticas</span>'
                '<span class="chip" style="color:' + fg_c + ';background:' + bg_c + ';">' + ("ATENCIÓN" if s["criticas_sesion"] else "OK") + '</span></div>'
                '<div class="kpi-num-big" style="color:#C62828">' + str(s["criticas_sesion"]) + '</div><div class="kpi-sub">R2L y U2R</div></div>',
                unsafe_allow_html=True)
    st.write("")

    izq, der = st.columns([1.2, 1])
    with izq:
        with st.container(key="card_sensor_feed"):
            filas = ""
            for e in reversed(s["feed"]):
                color = COLOR_CAT[e["cat"]]
                filas += ('<div class="feed-row" style="border-left-color:' + color + '"><span>#' + str(e["n"]) + ' · ' + e["hora"]
                          + ' · ' + html.escape(e["proto"]) + '</span><span style="color:' + color + ';font-weight:700;">'
                          + NOMBRES_AMENAZA[e["cat"]] + ' · ' + "{:.0f}%".format(e["certeza"] * 100) + '</span></div>')
            st.markdown('<div class="panel-title">Tráfico analizado en vivo</div>'
                        '<div class="panel-sub">Cada fila es un evento de red que el modelo clasificó</div>' + (filas or "Esperando eventos..."),
                        unsafe_allow_html=True)
    with der:
        with st.container(key="card_sensor_chart"):
            st.markdown('<div class="panel-title">Alertas por lote</div>'
                        '<div class="panel-sub">Las barras altas son ráfagas de ataque</div>', unsafe_allow_html=True)
            fig = go.Figure(go.Bar(x=list(range(1, len(s["serie"]) + 1)), y=s["serie"], marker_color="#E8843C"))
            fig.update_layout(height=230, margin=dict(l=10, r=10, t=5, b=10), paper_bgcolor="rgba(0,0,0,0)",
                              plot_bgcolor="rgba(0,0,0,0)", xaxis=dict(showgrid=False, title="Lote"),
                              yaxis=dict(gridcolor="#EEF1F5", title="Alertas"))
            st.plotly_chart(fig, width="stretch", config={"displayModeBar": False})

    with st.container(key="card_sensor_incs"):
        st.markdown('<div class="panel-title">Incidentes creados por el sensor</div>'
                    '<div class="panel-sub">Los más recientes de esta sesión. Quedan guardados en el Historial.</div>', unsafe_allow_html=True)
        recientes = listar_incidentes(usuario, analisis_id=s["sesion_id"])[:8]
        if not recientes:
            st.markdown('<div class="panel-sub">Todavía no hay incidentes: el tráfico va normal.</div>', unsafe_allow_html=True)
        else:
            filas = "".join('<tr><td class="mono">' + fmt_fecha(a["fecha"], corto=True) + '</td><td class="mono">#' + str(a["id"]) + '</td><td><b>'
                            + NOMBRES_AMENAZA[a["categoria"]] + '</b></td><td>' + chip(a["severidad"], SEV_ESTILO[a["severidad"]], True)
                            + '</td><td>' + chip(a["estado"], ESTADO_ESTILO[a["estado"]]) + '</td></tr>' for a in recientes)
            st.markdown('<div class="tabla-wrap"><table class="tabla-alertas"><thead><tr><th>HORA</th><th>ID</th><th>TIPO DE AMENAZA</th>'
                        '<th>SEVERIDAD</th><th>ESTADO</th></tr></thead><tbody>' + filas + '</tbody></table></div>', unsafe_allow_html=True)
            if st.button("Ver todos en el Historial ›", key="btn_sensor_hist", type="tertiary"):
                ir_seccion("historial")
                st.rerun()

    with st.expander("Evaluación de esta sesión (posible porque el simulador conoce la verdad)"):
        filas = ""
        for c in CLASES:
            v = s["verdad"][c]
            if c == "Normal":
                detalle_txt = str(v["alertados"]) + " falsa(s) alarma(s)"
                pct = "-"
            else:
                pct = "{:.0f}%".format(v["alertados"] / v["enviados"] * 100) if v["enviados"] else "-"
                detalle_txt = str(v["alertados"]) + " detectados · " + str(v["exactos"]) + " con el tipo correcto"
            filas += '<tr><td><b>' + NOMBRES_AMENAZA[c] + '</b></td><td>' + str(v["enviados"]) + '</td><td>' + detalle_txt + '</td><td>' + pct + '</td></tr>'
        st.markdown('<table class="tabla-alertas"><thead><tr><th>TIPO REAL</th><th>ENVIADOS</th><th>RESULTADO</th><th>% DETECTADO</th></tr></thead><tbody>'
                    + filas + '</tbody></table><div class="panel-sub" style="margin-top:8px;">El modelo detecta bien DoS y Probe; '
                    'R2L y U2R son su punto débil (limitación documentada del dataset NSL-KDD).</div>', unsafe_allow_html=True)


def vista_monitoreo(usuario):
    st.markdown('<div class="dash-title">◉ Monitoreo en vivo</div>'
                '<div class="dash-sub">Un sensor virtual envía tráfico de red simulado y Centinela lo analiza solo: clasifica cada evento, '
                'crea los incidentes y avisa. Funciona mientras esta pantalla esté abierta.</div><br>', unsafe_allow_html=True)
    if not os.path.exists(ARCHIVO_SENSOR):
        st.error("Falta el archivo " + ARCHIVO_SENSOR + " en la carpeta del proyecto. Cópialo junto a app.py y recarga.")
        return
    grupos = cargar_grupos_sensor()
    if "sensor" not in st.session_state:
        st.session_state["sensor"] = nuevo_estado_sensor()
    s = st.session_state["sensor"]

    with st.container(key="card_sensor_ctrl"):
        c1, c2, c3, c4 = st.columns([1.3, 1, 1, 1.8], vertical_alignment="bottom")
        with c1:
            if s["activo"]:
                st.button("■ Detener monitoreo", key="btn_sensor_toggle", on_click=detener_sensor, width="stretch")
            else:
                st.button("▶ Iniciar monitoreo", key="btn_sensor_toggle", type="primary", on_click=iniciar_sensor,
                          args=(usuario["id"],), width="stretch")
        with c2:
            st.select_slider("Intensidad de ataques", ["Baja", "Media", "Alta"], value="Media", key="sensor_intensidad")
        with c3:
            st.select_slider("Velocidad del tráfico", ["Lenta", "Media", "Rápida"], value="Media", key="sensor_ritmo")
        with c4:
            st.multiselect("Ataques que puede enviar", CATEGORIAS, default=CATEGORIAS, key="sensor_tipos")

    fragmento = st.fragment(run_every="2s" if s["activo"] else None)(panel_sensor)
    fragmento(usuario, grupos)


# =====================================================================
# PANTALLA: LOGIN
# =====================================================================
chat_flotante()

if pagina == "login":
    encabezado_tarjeta("CENTINELA", "Prototipo académico de monitoreo y análisis de<br>seguridad para pymes.")

    if st.session_state.pop("aviso_registro", False):
        st.success("Solicitud enviada. Cuando el administrador la apruebe, podrás entrar con tu correo y contraseña.")

    st.markdown('<div class="field-label">Correo corporativo</div>', unsafe_allow_html=True)
    correo = st.text_input("correo_input", placeholder="ejemplo@empresa.com", label_visibility="collapsed", key="correo")
    st.markdown('<div class="field-label">Contraseña</div>', unsafe_allow_html=True)
    clave = st.text_input("clave_input", type="password", placeholder="••••••••••", label_visibility="collapsed", key="clave")

    c1, c2 = st.columns([1, 1], vertical_alignment="center")
    with c1:
        st.checkbox("Recordar")
    with c2:
        if st.button("¿Olvidaste?", key="forgot", type="tertiary", width="stretch"):
            st.toast("Para restablecer tu contraseña, contacta al administrador de Centinela.")

    if st.button("Iniciar sesión →", key="btn_login", type="primary", width="stretch"):
        estado, usuario = autenticar(correo, clave)
        if estado == "aprobado":
            st.session_state["usuario"] = usuario
            for k in ("resultado", "incidente_abierto", "reporte", "sensor", "enlaces_res", "enlaces_hist", "enlaces_txt", "correos_res", "correo_txt", "buzon_cfg", "buzon_log", "buzon_stats", "buzon_ultimo_uid", "buzon_vigilando", "buzon_error", "buzon_clave_txt", "buzon_usuario", "buzon_ultimo_chequeo", "buzon_aviso_estado", "buzon_destino", "buzon_avisar", "buzon_nivel_aviso"):
                st.session_state.pop(k, None)
            for k in [k for k in st.session_state.keys() if str(k).startswith(("chat_", "cfg_", "pref_"))]:
                st.session_state.pop(k, None)
            st.session_state["seccion"] = "dashboard"
            st.session_state["pagina"] = "dashboard"
            if usuario["rol"] == "admin":
                n = contar_pendientes()
                if n:
                    st.session_state["toast"] = f"Tienes {n} solicitud(es) de registro pendiente(s)."
            st.rerun()
        elif estado == "pendiente":
            st.info("Tu solicitud sigue en revisión. El administrador aún no la aprueba. Si tienes dudas, usa el botón «Chatea con Centinela».")
        elif estado == "rechazado":
            st.error("Tu solicitud de registro fue rechazada. Contacta al administrador.")
        else:
            st.error("Correo o contraseña incorrectos.")

    st.markdown('<div class="divider-text">O CONTINÚA CON</div>', unsafe_allow_html=True)
    with st.container(key="footer_reg"):
        f1, f2 = st.columns([3, 2], vertical_alignment="center")
        with f1:
            st.markdown("¿Nueva organización?")
        with f2:
            st.button("Registrar empresa", key="ir_registro", type="tertiary", on_click=ir_a, args=("registro",))
    pie_tarjeta()

# =====================================================================
# PANTALLA: REGISTRO
# =====================================================================
elif pagina == "registro":
    encabezado_tarjeta("Registrar empresa", "Envía tu solicitud. Un administrador de Centinela<br>la revisará antes de darte acceso.")

    st.markdown('<div class="field-label">Nombre de la empresa</div>', unsafe_allow_html=True)
    r_empresa = st.text_input("reg_empresa", placeholder="Ej. Distribuidora Andina S.A.S.", label_visibility="collapsed", key="r_empresa")
    st.markdown('<div class="field-label">Correo corporativo</div>', unsafe_allow_html=True)
    r_correo = st.text_input("reg_correo", placeholder="ejemplo@empresa.com", label_visibility="collapsed", key="r_correo")
    st.markdown('<div class="field-label">Contraseña (mínimo 8 caracteres)</div>', unsafe_allow_html=True)
    r_clave = st.text_input("reg_clave", type="password", placeholder="••••••••••", label_visibility="collapsed", key="r_clave")
    st.markdown('<div class="field-label">Confirmar contraseña</div>', unsafe_allow_html=True)
    r_clave2 = st.text_input("reg_clave2", type="password", placeholder="••••••••••", label_visibility="collapsed", key="r_clave2")

    st.write("")
    if st.button("Enviar solicitud →", key="btn_registrar", type="primary", width="stretch"):
        ok, msg = crear_solicitud(r_empresa, r_correo, r_clave, r_clave2)
        if ok:
            for k in ("r_empresa", "r_correo", "r_clave", "r_clave2"):
                st.session_state.pop(k, None)
            st.session_state["aviso_registro"] = True
            st.session_state["pagina"] = "login"
            st.rerun()
        else:
            st.error(msg)

    st.markdown('<div class="divider-text">¿YA TIENES CUENTA?</div>', unsafe_allow_html=True)
    with st.container(key="footer_reg"):
        g1, g2 = st.columns([3, 2], vertical_alignment="center")
        with g1:
            st.markdown("Ya estoy registrado")
        with g2:
            st.button("Iniciar sesión", key="ir_login", type="tertiary", on_click=ir_a, args=("login",))
    pie_tarjeta()

# =====================================================================
# PANTALLA: DASHBOARD
# =====================================================================
elif pagina == "dashboard":
    usuario = st.session_state["usuario"]
    es_admin = usuario["rol"] == "admin"
    empresa_raw = "Todas las empresas" if es_admin else usuario["empresa"]
    empresa = html.escape(usuario["empresa"])
    correo_u = html.escape(usuario["correo"])
    nombre_raw = usuario["correo"].split("@")[0].replace(".", " ").title()
    nombre_corto = html.escape(nombre_raw)
    iniciales = "".join(p[0].upper() for p in usuario["correo"].split("@")[0].replace(".", " ").split()[:2]) or "CE"
    pendientes = contar_pendientes() if es_admin else 0
    nuevos = contar_nuevos(usuario)

    seccion = st.session_state.get("seccion", "dashboard")
    if seccion == "solicitudes" and not es_admin:
        seccion = "dashboard"

    NAV = [("dashboard", "▦", "Dashboard"), ("correos", "◆", "Analizar correos"), ("conector", "⇄", "Conector de correo"), ("monitoreo", "◉", "Monitoreo en vivo"), ("enlaces", "◈", "Verificar enlaces"),
           ("solicitudes", "✉", "Solicitudes"), ("historial", "↻", "Historial"), ("reportes", "▤", "Reportes"), ("config", "⚙", "Configuración")]

    with st.sidebar:
        st.markdown(
            f'<div class="sb-logo"><div class="box"><img src="data:image/png;base64,{logo_b64}"></div>'
            f'<div><b>CENTINELA</b><small>CIBERSEGURIDAD PYME</small></div></div>', unsafe_allow_html=True)
        for clave_nav, icono, nombre in NAV:
            if clave_nav == "solicitudes" and not es_admin:
                continue
            etiqueta = icono + "  " + nombre
            if clave_nav == "solicitudes" and pendientes:
                etiqueta += " (" + str(pendientes) + ")"
            if clave_nav == "monitoreo" and st.session_state.get("sensor", {}).get("activo"):
                etiqueta += "  ●"
            if clave_nav == "historial" and nuevos:
                etiqueta += " (" + str(nuevos) + ")"
            st.button(etiqueta, key="nav_" + clave_nav, type="primary" if seccion == clave_nav else "tertiary",
                      on_click=ir_seccion, args=(clave_nav,), width="stretch")
        st.markdown(
            f'<div class="sb-user"><div class="avatar">{iniciales}</div>'
            f'<div><div class="nm">{nombre_corto}</div><div class="rl">{"Administrador" if es_admin else "Usuario TI"}</div></div></div>',
            unsafe_allow_html=True)
        st.write("")
        if st.button("Cerrar sesión", key="btn_logout", width="stretch"):
            for k in ("usuario", "resultado", "seccion", "incidente_abierto", "reporte", "sensor", "enlaces_res", "enlaces_hist", "enlaces_txt", "correos_res", "correo_txt", "buzon_cfg", "buzon_log", "buzon_stats", "buzon_ultimo_uid", "buzon_vigilando", "buzon_error", "buzon_clave_txt", "buzon_usuario", "buzon_ultimo_chequeo", "buzon_aviso_estado", "buzon_destino", "buzon_avisar", "buzon_nivel_aviso"):
                st.session_state.pop(k, None)
            for k in [k for k in st.session_state.keys() if str(k).startswith(("chat_", "cfg_", "pref_"))]:
                st.session_state.pop(k, None)
            st.session_state["pagina"] = "login"
            st.rerun()

    campana = pendientes if es_admin else nuevos
    campana_html = f'<span class="bell-badge">{campana}</span>' if campana else ""
    st.markdown(
        f'<div class="topbar-card"><div><div class="company-name">{empresa}</div>'
        f'<div class="company-sub">Prototipo académico · {correo_u}</div></div>'
        f'<div class="topbar-right"><span class="badge-live"><span class="dot"></span>Motor de detección activo</span>'
        f'<div class="bell"><svg width="22" height="22" viewBox="0 0 24 24" fill="none" stroke="#E0A93B" stroke-width="2" stroke-linecap="round" stroke-linejoin="round">'
        f'<path d="M18 8a6 6 0 0 0-12 0c0 7-3 9-3 9h18s-3-2-3-9"/><path d="M13.7 21a2 2 0 0 1-3.4 0"/></svg>{campana_html}</div>'
        f'<div class="user-chip"><div class="avatar">{iniciales}</div><span>{nombre_corto}</span></div></div></div>',
        unsafe_allow_html=True)

    toast_msg = st.session_state.pop("toast", None)
    if toast_msg:
        st.toast(toast_msg)

    # ------------------------- SOLICITUDES (solo admin) -------------------------
    if seccion == "solicitudes" and es_admin:
        st.markdown('<div class="dash-title">✉ Solicitudes de registro</div>'
                    '<div class="dash-sub">Aprueba o rechaza a las empresas que quieren entrar a Centinela.</div><br>',
                    unsafe_allow_html=True)
        pend = listar_usuarios("pendiente")
        if not pend:
            with st.container(key="card_sin_pend"):
                st.info("No hay solicitudes pendientes.")
        for u in pend:
            with st.container(key=f"card_sol_{u['id']}"):
                a, b, c = st.columns([4, 1, 1], vertical_alignment="center")
                with a:
                    correo_sol = html.escape(u["correo"]).replace("@", "&#64;")
                    st.markdown(f"**{html.escape(u['empresa'])}**  \n{correo_sol} · solicitó el {u['fecha']}")
                with b:
                    if st.button("Aprobar", key=f"ap_{u['id']}", type="primary", width="stretch"):
                        cambiar_estado(u["id"], "aprobado")
                        st.rerun()
                with c:
                    if st.button("Rechazar", key=f"re_{u['id']}", width="stretch"):
                        cambiar_estado(u["id"], "rechazado")
                        st.rerun()
        hist = listar_usuarios("aprobado") + listar_usuarios("rechazado")
        if hist:
            with st.container(key="card_hist_sol"):
                st.markdown('<div class="panel-title">Historial de decisiones</div>', unsafe_allow_html=True)
                df = pd.DataFrame(hist).rename(columns={"empresa": "Empresa", "correo": "Correo", "fecha": "Solicitud", "estado": "Estado"}).drop(columns=["id"])
                st.dataframe(df, width="stretch", hide_index=True)
        st.stop()

    # ------------------------- MONITOREO EN VIVO -------------------------
    if seccion == "monitoreo":
        vista_monitoreo(usuario)
        st.stop()

    # ------------------------- CONECTOR DE CORREO -------------------------
    if seccion == "conector":
        vista_conector(usuario)
        st.stop()

    # ------------------------- ANALIZAR CORREOS -------------------------
    if seccion == "correos":
        vista_correos(usuario)
        st.stop()

    # ------------------------- VERIFICADOR DE ENLACES -------------------------
    if seccion == "enlaces":
        vista_enlaces(usuario)
        st.stop()

    # ------------------------- HISTORIAL / INCIDENTE -------------------------
    if seccion == "historial":
        abierto = st.session_state.get("incidente_abierto")
        if abierto:
            vista_incidente(usuario, abierto, nombre_raw)
        else:
            vista_historial(usuario)
        st.stop()

    # ------------------------- REPORTES -------------------------
    if seccion == "reportes":
        vista_reportes(usuario, empresa_raw, nombre_raw)
        st.stop()

    if seccion == "config":
        vista_configuracion(usuario)
        st.stop()

    # ------------------------- DASHBOARD -------------------------
    if es_admin and pendientes:
        with st.container(key="card_banner"):
            b1, b2 = st.columns([5, 1], vertical_alignment="center")
            with b1:
                st.markdown(f"🔔 **{pendientes} solicitud(es) de registro** esperan tu aprobación.")
            with b2:
                st.button("Revisar", key="btn_revisar", type="primary", width="stretch",
                          on_click=ir_seccion, args=("solicitudes",))

    ultimo_f = ultimo_analisis(usuario)
    ultimo = fmt_fecha(ultimo_f) if ultimo_f else "—"
    st.markdown(
        f'<div class="dash-head"><div><div class="dash-title">🛡️ {empresa} <span class="tag-pyme">PYME</span></div>'
        f'<div class="dash-sub">Resultados del último análisis · Motor de Machine Learning (XGBoost)</div></div>'
        f'<div class="last-box"><div class="last-label">Último análisis</div><div class="last-val">{ultimo}</div></div></div>',
        unsafe_allow_html=True)

    # ---------- Subir CSV y monitorear ----------
    with st.container(key="card_subir"):
        st.markdown('<div class="panel-title">Subir archivo de eventos (CSV)</div>'
                    '<div class="panel-sub">El sistema clasifica cada evento con el modelo de IA y lo muestra en vivo. Cada análisis queda guardado en el Historial.</div>',
                    unsafe_allow_html=True)
        archivo = st.file_uploader("CSV", type=["csv"], label_visibility="collapsed")

        if archivo is not None:
            datos = pd.read_csv(archivo)
            faltantes = [c for c in COLUMNAS_MODELO if c not in datos.columns]
            if faltantes:
                st.error(f"Al archivo le faltan estas columnas: {faltantes}")
            else:
                ca, cb = st.columns(2)
                with ca:
                    velocidad = st.slider("Velocidad (segundos entre eventos)", 0.05, 1.0, 0.15)
                with cb:
                    tope = max(5, min(100, len(datos)))
                    max_eventos = st.slider("Cuántos eventos simular", 5, tope, min(30, tope))

                if st.button("▶ Iniciar monitoreo en vivo", key="btn_monitoreo", type="primary", width="stretch"):
                    max_eventos = min(max_eventos, len(datos))
                    X = datos[COLUMNAS_MODELO].copy()
                    for col, encoder in encoders.items():
                        X[col] = X[col].map(lambda v: v if v in encoder.classes_ else encoder.classes_[0])
                        X[col] = encoder.transform(X[col])

                    lista = st.container()
                    conteo = {"Normal": 0, "DoS": 0, "Probe": 0, "R2L": 0, "U2R": 0}
                    incidentes, serie = [], []

                    for i in range(max_eventos):
                        proba = modelo.predict_proba(X.iloc[[i]])[0]
                        idx = int(proba.argmax())
                        cat = target_encoder.inverse_transform([modelo.classes_[idx]])[0]
                        certeza = float(proba[idx])
                        conteo[cat] += 1
                        serie.append(1 if cat != "Normal" else 0)
                        color = COLOR_CAT[cat]
                        if cat != "Normal":
                            fila_dato = datos.iloc[i]
                            det = {c: (fila_dato[c].item() if hasattr(fila_dato[c], "item") else fila_dato[c]) for c in COLUMNAS_DETALLE}
                            incidentes.append({
                                "fecha": ahora_iso(), "categoria": cat, "severidad": SEVERIDAD[cat],
                                "activo": ACTIVO_SIMULADO[cat], "protocolo": str(fila_dato["protocol_type"]),
                                "evento_num": i + 1, "certeza": certeza, "detalle": json.dumps(det, default=str),
                            })
                        with lista:
                            st.markdown(
                                f'<div class="evento-row" style="border-left-color:{color}">'
                                f'<span>Evento #{i+1} · {html.escape(str(datos.iloc[i]["protocol_type"]))}</span>'
                                f'<span style="color:{color}; font-weight:700;">{NOMBRES_AMENAZA[cat]}</span></div>',
                                unsafe_allow_html=True)
                        time.sleep(velocidad)

                    aid, ts = guardar_analisis(usuario["id"], conteo, max_eventos, incidentes)
                    st.session_state["resultado"] = {
                        "analisis_id": aid, "conteo": conteo, "total": max_eventos, "serie": serie, "fecha": fmt_fecha(ts),
                    }
                    st.rerun()

    # ---------- Resultados ----------
    resultado = st.session_state.get("resultado")
    if resultado:
        conteo, total, serie = resultado["conteo"], resultado["total"], resultado["serie"]
        incs = listar_incidentes(usuario, analisis_id=resultado["analisis_id"])
        no_normales = total - conteo["Normal"]
        score = round(no_normales / total * 100) if total else 0
        criticos = sum(1 for a in incs if a["severidad"] == "CRÍTICO")
        pct_normal = round(conteo["Normal"] / total * 100, 1) if total else 0
        if score < 33:
            nivel = "BAJO"
        elif score < 66:
            nivel = "MEDIO"
        else:
            nivel = "ALTO"
        fg_n, bg_n = SEV_ESTILO[nivel]
        color_barra = {"BAJO": "#3C8C40", "MEDIO": "#E0A93B", "ALTO": "#E23D3D"}[nivel]

        k1, k2, k3, k4 = st.columns(4)
        with k1:
            with st.container(key="card_kpi_gauge"):
                st.markdown(
                    f'<div class="kpi-top"><span class="kpi-label-top">Nivel de amenaza global</span>'
                    f'<span class="chip" style="color:{fg_n};background:{bg_n};">{nivel}</span></div>', unsafe_allow_html=True)
                fig_g = go.Figure(go.Indicator(
                    mode="gauge+number", value=score, number={"suffix": "/100", "font": {"size": 26, "color": "#16202B"}},
                    gauge={"axis": {"range": [0, 100], "visible": False}, "bar": {"color": color_barra, "thickness": 0.35},
                           "bgcolor": "#EEF1F5", "borderwidth": 0}))
                fig_g.update_layout(height=105, margin=dict(l=10, r=10, t=8, b=0), paper_bgcolor="rgba(0,0,0,0)")
                st.plotly_chart(fig_g, width="stretch", config={"displayModeBar": False})
                st.markdown('<div class="kpi-sub" style="text-align:center;margin-top:-6px;">Riesgo estimado por el modelo</div>', unsafe_allow_html=True)
        with k2:
            st.markdown(f'<div class="kpi-card"><div class="kpi-label-top">Eventos analizados</div>'
                        f'<div class="kpi-num-big">{total}</div>'
                        f'<div class="kpi-sub">último lote · <b>{pct_normal}%</b> tráfico normal</div></div>', unsafe_allow_html=True)
        with k3:
            st.markdown(f'<div class="kpi-card"><div class="kpi-label-top">Alertas detectadas</div>'
                        f'<div class="kpi-num-big">{len(incs)}</div>'
                        f'<div class="kpi-sub">{no_normales} de {total} eventos no fueron normales</div></div>', unsafe_allow_html=True)
        with k4:
            fg_c, bg_c = SEV_ESTILO["CRÍTICO"]
            st.markdown(f'<div class="kpi-card"><div class="kpi-top"><span class="kpi-label-top">Críticos</span>'
                        f'<span class="chip" style="color:{fg_c};background:{bg_c};">{"ATENCIÓN" if criticos else "OK"}</span></div>'
                        f'<div class="kpi-num-big" style="color:#C62828">{criticos}</div>'
                        f'<div class="kpi-sub">Vectores: R2L y U2R</div></div>', unsafe_allow_html=True)

        st.write("")
        g1, g2 = st.columns([1.55, 1])
        with g1:
            with st.container(key="card_linea"):
                st.markdown('<div class="panel-title">Evolución de las amenazas detectadas</div>'
                            '<div class="panel-sub">Amenazas acumuladas según el orden de llegada de los eventos</div>', unsafe_allow_html=True)
                acum = pd.Series(serie).cumsum()
                fig_l = go.Figure(go.Scatter(x=list(range(1, len(acum) + 1)), y=acum, mode="lines",
                                             line=dict(color="#2E7CE4", width=3, shape="spline"),
                                             fill="tozeroy", fillcolor="rgba(46,124,228,0.12)"))
                fig_l.update_layout(height=215, margin=dict(l=10, r=10, t=10, b=10), paper_bgcolor="rgba(0,0,0,0)",
                                    plot_bgcolor="rgba(0,0,0,0)", xaxis=dict(showgrid=False, title="Evento #"),
                                    yaxis=dict(gridcolor="#EEF1F5", title="Acumuladas"))
                st.plotly_chart(fig_l, width="stretch", config={"displayModeBar": False})
        with g2:
            with st.container(key="card_donut"):
                st.markdown('<div class="panel-title">Eventos por categoría</div>'
                            '<div class="panel-sub">Clasificación del modelo · NSL-KDD</div>', unsafe_allow_html=True)
                d1, d2 = st.columns([1, 1.15])
                with d1:
                    st.plotly_chart(figura_donut(conteo, total), width="stretch", config={"displayModeBar": False})
                with d2:
                    st.markdown(leyenda_html(conteo, total), unsafe_allow_html=True)

        with st.container(key="card_tabla"):
            st.markdown('<div class="panel-title">Alertas detectadas</div>'
                        '<div class="panel-sub">Incidentes priorizados por severidad · la hora es la del análisis; el activo afectado es simulado (el dataset no trae equipos)</div>',
                        unsafe_allow_html=True)
            if not incs:
                st.info("No se detectaron amenazas en los eventos analizados.")
            else:
                orden = {"CRÍTICO": 0, "ALTO": 1, "MEDIO": 2, "BAJO": 3}
                filas = ""
                for a in sorted(incs, key=lambda x: (orden[x["severidad"]], -x["id"])):
                    filas += ('<tr><td class="mono">' + fmt_fecha(a["fecha"], corto=True) + '</td><td class="mono">#' + str(a["id"]) + '</td>'
                              '<td><b>' + NOMBRES_AMENAZA[a["categoria"]] + '</b></td><td>' + html.escape(a["activo"]) + '</td>'
                              '<td>' + chip(a["severidad"], SEV_ESTILO[a["severidad"]], True) + '</td>'
                              '<td>' + chip(a["estado"], ESTADO_ESTILO[a["estado"]]) + '</td></tr>')
                st.markdown(
                    '<div class="tabla-wrap"><table class="tabla-alertas"><thead><tr><th>HORA</th><th>ID</th><th>TIPO DE AMENAZA</th>'
                    '<th>ACTIVO (SIMULADO)</th><th>SEVERIDAD</th><th>ESTADO</th></tr></thead><tbody>' + filas + '</tbody></table></div>',
                    unsafe_allow_html=True)
                st.button("Ver todas las alertas en el Historial ›", key="btn_ver_todas", type="tertiary",
                          on_click=ir_seccion, args=("historial",))
